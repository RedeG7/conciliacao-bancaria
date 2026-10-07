#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Interface web (Streamlit) para o script conciliacao_bancaria.py

Rodar:
    streamlit run app_conciliacao.py

Permite:
  - Upload direto dos 3 arquivos (extrato, razao, balancete), OU
  - Informar o caminho de uma pasta no disco e o app identifica os
    arquivos automaticamente pelo nome (extrato/fluxo, razao, balancete).
  - Informar codigo da empresa no Dominio + demais parametros.
  - Ver o espelho de conciliacao na tela e baixar os 3 entregaveis
    (espelho .md, memoria .csv, importacao Dominio .txt).
"""

from __future__ import annotations

import io
import os
import re
import shutil
import tempfile
import zipfile
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Optional

import streamlit as st

import auth
import clientes
import historico
import conciliacao_bancaria as cb
from rpa import core as rpa_core
from rpa import registry as rpa_registry
from rpa.sefazgo_nfe import arquivos as nfgo_arquivos

st.set_page_config(page_title="Hub App", page_icon="🧩", layout="wide")

# O menu "⋮" (tema/Print/Record screen) e o botao "Deploy" no canto
# superior direito sao texto fixo em ingles, sem suporte a traducao -
# escondidos por isso. IMPORTANTE: nao esconder [data-testid='stToolbar']
# inteiro - o botao de EXPANDIR a sidebar quando ela esta minimizada
# (stExpandSidebarButton) mora dentro desse mesmo container, e escondendo
# o pai ele some junto, prendendo o usuario com a sidebar fechada sem
# jeito de reabrir.
st.markdown(
    "<style>#MainMenu, [data-testid='stMainMenu'], [data-testid='stMainMenuButton'],"
    "[data-testid='stBaseButton-header']"
    "{visibility:hidden; display:none;}"
    # o aviso "Press Enter to submit form"/"Press Enter to apply" que o
    # Streamlit mostra ao digitar num campo tambem e texto fixo em ingles,
    # sem opcao de traducao - escondido pelo mesmo motivo.
    "[data-testid='InputInstructions']{visibility:hidden; display:none;}</style>",
    unsafe_allow_html=True,
)

# Usuarios/escritorios/clientes/historico ficam no Postgres (DATABASE_URL) -
# ver auth.py. Isso evita corrida de escrita entre admins de escritorios
# diferentes mexendo ao mesmo tempo (cada operacao e uma transacao atomica
# no banco).

if "DATABASE_URL" not in os.environ:
    # roda esse arquivo direto (sem docker compose) sem apontar pra um
    # Postgres - o mais comum e alguem tentando ACESSAR o sistema em vez de
    # rodar em modo desenvolvedor: o site de producao ja fica no ar em
    # https://hub.redeg7.com, nao precisa (nem deveria) rodar isso localmente
    # pra usar o sistema.
    st.error(
        "⚠️ **Este app não está configurado com um banco de dados** "
        "(variável de ambiente `DATABASE_URL` não definida)."
    )
    st.info(
        "**Se você quer só usar o sistema:** acesse "
        "[hub.redeg7.com](https://hub.redeg7.com) — não precisa rodar nada "
        "no seu computador.\n\n"
        "**Se você é desenvolvedor e quer testar localmente:** defina a "
        "variável `DATABASE_URL` apontando pra um Postgres (local ou via "
        "túnel) antes de rodar `streamlit run app_conciliacao.py`, ex.:\n\n"
        "```\nDATABASE_URL=postgresql://usuario:senha@localhost:5432/conciliacao "
        "streamlit run app_conciliacao.py\n```"
    )
    st.stop()


@st.cache_resource
def _garantir_schema_extra() -> bool:
    # auth primeiro: clientes/historico tem FK pra escritorios(id).
    auth.garantir_schema()
    clientes.garantir_schema()
    historico.garantir_schema()
    rpa_core.garantir_schema()
    return True


try:
    _garantir_schema_extra()
except Exception as exc:
    st.error(
        f"⚠️ **Não consegui conectar ao banco de dados configurado.** "
        f"Confira se o Postgres em `DATABASE_URL` está no ar e acessível.\n\n"
        f"Detalhe técnico: `{exc}`"
    )
    st.stop()

# ---------------------------------------------------------------------------
# Login / autenticacao multi-escritorio - nada do app roda sem sessao
# autenticada. Varios escritorios de contabilidade usam o mesmo app ao
# mesmo tempo, cada um enxergando so os proprios usuarios (isolamento por
# escritorio_id). Ha sempre um super_admin_global (bootstrap "admin"/
# "admin123" no primeiro uso, escritorio "redeg7") que enxerga todos os
# escritorios; cada escritorio tambem tem seu proprio admin_escritorio,
# capaz de redefinir a senha de quem esqueceu sem depender do global.
# ---------------------------------------------------------------------------

_CHAVES_SESSAO_LOGIN = (
    "usuario_logado", "papel_usuario", "nome_usuario", "escritorio_id",
    "escritorio_nome", "deve_trocar_senha", "mostrar_trocar_senha",
)


_COOKIE_SESSAO = "sessao_token"


def _https_ativo() -> bool:
    try:
        return str(st.context.url or "").startswith("https://")
    except Exception:
        return False


def _definir_cookie_sessao(token: str, dias: int) -> None:
    """Grava o token de sessao persistente no cookie do navegador (via um
    scriptzinho injetado - o Streamlit nao expoe um jeito nativo de setar
    cookie a partir do Python). Nao e HttpOnly (o Streamlit nao da controle
    sobre os headers HTTP de resposta pra isso) - fica no mesmo nivel de
    protecao de localStorage, aceitavel aqui porque o app nao renderiza
    HTML/JS vindo de outro usuario (sem superficie de XSS entre tenants)."""
    seguro = "; Secure" if _https_ativo() else ""
    st.components.v1.html(
        f"<script>document.cookie = "
        f"'{_COOKIE_SESSAO}={token}; path=/; max-age={dias * 86400}; SameSite=Lax{seguro}';</script>",
        height=0,
    )


def _limpar_cookie_sessao() -> None:
    st.components.v1.html(
        f"<script>document.cookie = '{_COOKIE_SESSAO}=; path=/; max-age=0; SameSite=Lax';</script>",
        height=0,
    )


def _agendar_cookie_sessao(acao: str, valor: str = "") -> None:
    """`st.rerun()` interrompe a execucao do script na hora - se o
    componente que grava/apaga o cookie fosse renderizado bem antes de um
    rerun(), o navegador nunca chegaria a rodar aquele JS (a tela e trocada
    antes). Por isso o pedido de gravar/apagar cookie fica "agendado" no
    session_state e so e efetivamente renderizado no INICIO do proximo
    script run (via `_renderizar_cookie_pendente`), depois que o rerun ja
    aconteceu - mesmo motivo/padrao do `_flash()` pra mensagens de sucesso."""
    st.session_state["_cookie_pendente"] = (acao, valor)


def _renderizar_cookie_pendente() -> None:
    pendente = st.session_state.pop("_cookie_pendente", None)
    if not pendente:
        return
    acao, valor = pendente
    if acao == "set":
        _definir_cookie_sessao(valor, auth.DURACAO_SESSAO_DIAS)
    elif acao == "clear":
        _limpar_cookie_sessao()


def _popular_sessao(dados: dict, token: Optional[str] = None) -> None:
    escritorios = auth.carregar_escritorios()
    eid = dados.get("escritorio_id", "")
    st.session_state["usuario_logado"] = dados["usuario"]
    st.session_state["papel_usuario"] = dados.get("papel", auth.PAPEL_USUARIO)
    st.session_state["nome_usuario"] = dados.get("nome", dados["usuario"])
    st.session_state["escritorio_id"] = eid
    st.session_state["escritorio_nome"] = escritorios.get(eid, {}).get("nome", eid)
    st.session_state["deve_trocar_senha"] = bool(dados.get("deve_trocar_senha"))
    if token:
        st.session_state[_COOKIE_SESSAO] = token


def _resetar_estado_nao_login() -> None:
    """Limpa TUDO do session_state exceto as chaves de login/cookie (ver
    _CHAVES_SESSAO_LOGIN) - em especial apaga do disco a pasta temporaria
    de uploads (tmpdir) e limpa o @st.cache_data de processamento.

    CAUSA RAIZ de um bug de vazamento entre escritorios confirmado pelo
    usuario: st.session_state sobrevive a troca de usuario quando a MESMA
    aba do navegador faz logout de um escritorio e login de outro (nao e
    uma conexao nova) - sem essa limpeza, o extrato/razao/balancete que o
    escritorio A tinha enviado (pasta e nomes fixos: extrato.*, razao.*,
    balancete.* dentro de st.session_state.tmpdir) continuava no disco e
    podia ser reaproveitado pelo proximo login naquela aba, fazendo o
    escritorio B "ler" arquivo/resultado de outro escritorio mesmo depois
    de enviar a propria planilha. Chamada tanto no logout quanto logo
    antes de popular uma sessao nova por login, como cinto-e-suspensorio."""
    tmpdir = st.session_state.get("tmpdir")
    if tmpdir:
        shutil.rmtree(tmpdir, ignore_errors=True)
    for chave in list(st.session_state.keys()):
        if chave not in _CHAVES_SESSAO_LOGIN and chave != _COOKIE_SESSAO:
            st.session_state.pop(chave, None)
    st.cache_data.clear()


def _fazer_logout() -> None:
    auth.remover_sessao(st.session_state.get(_COOKIE_SESSAO, ""))
    _resetar_estado_nao_login()
    for chave in _CHAVES_SESSAO_LOGIN:
        st.session_state.pop(chave, None)
    st.session_state.pop(_COOKIE_SESSAO, None)
    _agendar_cookie_sessao("clear")
    st.rerun()


def _estilo_hub_app() -> None:
    """CSS + fundo decorativo compartilhados pelas telas de autenticacao
    (login e trocar senha) - a identidade visual do Hub App (fundo escuro,
    ondas e icone de quebra-cabeca em degrade verde). Cada tela monta seu
    proprio cabecalho/conteudo em cima disso com as classes happ-*."""
    st.markdown(
        """
        <style>
        #MainMenu, header[data-testid="stHeader"], footer {visibility:hidden; height:0;}
        .block-container {padding-top:1.5rem !important; padding-bottom:2rem !important; max-width:1400px !important;}
        .stApp {
            background:
                radial-gradient(1100px 600px at 85% -10%, rgba(34,224,138,0.16), transparent 60%),
                radial-gradient(900px 500px at -10% 110%, rgba(34,224,138,0.10), transparent 60%),
                linear-gradient(160deg, #070b14 0%, #0b1220 55%, #060a12 100%);
        }
        .happ-bg-decor {position:fixed; inset:0; z-index:0; overflow:hidden; pointer-events:none;}
        .happ-header, .happ-dept-list, .happ-side-tagline, .happ-features, .happ-footer,
        div[data-testid="stForm"] {position:relative; z-index:1;}
        .happ-bg-decor .puzzle {
            position:absolute; top:-40px; right:-40px; font-size:420px; line-height:1;
            opacity:0.10; transform:rotate(8deg); filter:hue-rotate(70deg) saturate(1.6) grayscale(0.15);
        }
        .happ-bg-decor svg {position:absolute; left:0; bottom:-40px; width:100%; opacity:0.55;}

        .happ-header {text-align:center; margin-bottom:0.2rem;}
        .happ-logo-row {display:flex; align-items:center; justify-content:center; gap:14px; margin-bottom:2px;}
        .happ-logo-icon {
            font-size:44px; line-height:1;
            filter:hue-rotate(70deg) saturate(1.6) brightness(1.15) drop-shadow(0 0 12px rgba(34,224,138,.5));
        }
        .happ-logo-text {font-size:42px; font-weight:800; letter-spacing:-1px; color:#f5f7fa;}
        .happ-logo-text .happ-app {
            background:linear-gradient(90deg,#22e08a,#a8e63d);
            -webkit-background-clip:text; background-clip:text; color:transparent;
        }
        .happ-welcome {font-size:26px; font-weight:800; margin:10px 0 4px; color:#f5f7fa;}
        .happ-subtitle {font-size:15.5px; color:#93a1b7; margin:0 0 8px;}
        .happ-tagline {font-size:13px; color:#6f7f97;}
        .happ-tagline b {color:#9fb0c8; font-weight:600;}

        .happ-dept-list {display:flex; flex-direction:column; gap:18px; padding-top:26px;}
        .happ-dept-item {display:flex; align-items:center; gap:12px;}
        .happ-dept-icon {
            width:42px; height:42px; border-radius:11px; flex:none;
            display:flex; align-items:center; justify-content:center;
            background:rgba(255,255,255,0.035); border:1px solid rgba(255,255,255,0.09); font-size:18px;
        }
        .happ-dept-name {font-weight:700; font-size:14px; color:#f5f7fa;}
        .happ-dept-desc {font-size:11.5px; color:#93a1b7;}

        div[data-testid="stForm"] {
            background:rgba(255,255,255,0.035); border:1px solid rgba(255,255,255,0.09);
            border-radius:18px; padding:1.8rem 2rem 1.4rem;
        }
        div[data-testid="stForm"] label p {font-weight:700; color:#f5f7fa; font-size:14px;}
        div[data-testid="stTextInput"] {position:relative;}
        div[data-testid="stTextInput"] input {
            background:#111a2c !important; border:1px solid rgba(255,255,255,0.10) !important;
            border-radius:10px !important; color:#f5f7fa !important; padding-left:42px !important;
        }
        div[data-testid="stTextInput"]:has(input[aria-label="Usuário"])::before {
            content:"👤"; position:absolute; left:14px; top:41px; font-size:14px; opacity:.6; z-index:5;
        }
        div[data-testid="stTextInput"]:has(input[type="password"])::before {
            content:"🔒"; position:absolute; left:14px; top:41px; font-size:14px; opacity:.6; z-index:5;
        }
        .happ-remember-row {
            display:flex; align-items:center; justify-content:space-between;
            margin:2px 0 18px; font-size:12.5px; color:#93a1b7;
        }
        .happ-remember-row span.happ-forgot {color:#22e08a; font-weight:600;}
        div[data-testid="stFormSubmitButton"] button {
            width:100%; background:linear-gradient(90deg,#22e08a,#a8e63d) !important; color:#06210f !important;
            font-weight:800 !important; border:none !important; border-radius:10px !important; padding:0.75rem !important;
        }
        div[data-testid="stButton"] button {
            background:transparent !important; color:#93a1b7 !important;
            border:1px solid rgba(255,255,255,0.14) !important;
        }

        .happ-side-tagline {padding-top:70px;}
        .happ-side-tagline p {font-style:italic; font-size:19px; line-height:1.35; color:#dfe6ef; font-weight:500; margin:0 0 12px;}
        .happ-underline {width:60px; height:4px; border-radius:4px; background:linear-gradient(90deg,#22e08a,#a8e63d);}

        .happ-features {display:flex; justify-content:center; gap:70px; margin-top:44px; flex-wrap:wrap;}
        .happ-feature {display:flex; align-items:center; gap:12px;}
        .happ-feature .ficon {font-size:20px; color:#22e08a;}
        .happ-feature .ftitle {font-weight:700; font-size:13.5px; color:#f5f7fa;}
        .happ-feature .fdesc {font-size:11.5px; color:#93a1b7;}

        .happ-footer {text-align:center; margin-top:26px; font-size:11px; letter-spacing:2px; color:#4a5773;}
        </style>

        <div class="happ-bg-decor">
            <div class="puzzle">🧩</div>
            <svg viewBox="0 0 1600 260" preserveAspectRatio="none">
                <path d="M0,140 C300,220 500,60 850,120 C1150,175 1350,80 1600,140 L1600,260 L0,260 Z" fill="rgba(34,224,138,0.10)"/>
                <path d="M0,180 C320,240 620,120 900,160 C1200,200 1380,140 1600,190 L1600,260 L0,260 Z" fill="rgba(34,224,138,0.16)"/>
            </svg>
        </div>
        """,
        unsafe_allow_html=True,
    )


def _tela_login() -> None:
    """Tela de login com a identidade visual do Hub App: lista decorativa
    de departamentos e cartao de login. "Lembrar de mim" e "Esqueceu a
    senha?" sao so visuais - o sistema ainda nao tem essas funcoes (a
    sessao ja fica lembrada via cookie, e reset de senha e feito por um
    admin em Gerenciar Usuarios)."""
    _estilo_hub_app()
    st.markdown(
        """
        <div class="happ-header">
            <div class="happ-logo-row">
                <div class="happ-logo-icon">🧩</div>
                <div class="happ-logo-text">Hub <span class="happ-app">APP</span></div>
            </div>
            <div class="happ-welcome">Bem-vindo!</div>
            <div class="happ-subtitle">Sua central de automação contábil para todos os departamentos.</div>
            <div class="happ-tagline"><b>Mais eficiência</b> &nbsp;|&nbsp; <b>Mais integração</b> &nbsp;|&nbsp; <b>Mais resultados</b></div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    col_esq, col_meio, col_dir = st.columns([1, 1.3, 1])

    with col_esq:
        st.markdown(
            """
            <div class="happ-dept-list">
                <div class="happ-dept-item"><div class="happ-dept-icon">👥</div>
                    <div><div class="happ-dept-name">DP</div><div class="happ-dept-desc">Gestão de Pessoas</div></div></div>
                <div class="happ-dept-item"><div class="happ-dept-icon">📄</div>
                    <div><div class="happ-dept-name">Fiscal</div><div class="happ-dept-desc">Apuração e Obrigações</div></div></div>
                <div class="happ-dept-item"><div class="happ-dept-icon">📈</div>
                    <div><div class="happ-dept-name">Contábil</div><div class="happ-dept-desc">Demonstrativos e Relatórios</div></div></div>
                <div class="happ-dept-item"><div class="happ-dept-icon">⚙️</div>
                    <div><div class="happ-dept-name">Societário</div><div class="happ-dept-desc">Processos e Documentos</div></div></div>
                <div class="happ-dept-item"><div class="happ-dept-icon">☁️</div>
                    <div><div class="happ-dept-name">TI</div><div class="happ-dept-desc">Integrações e Automação</div></div></div>
            </div>
            """,
            unsafe_allow_html=True,
        )

    with col_meio:
        with st.form("login_form"):
            usuario = st.text_input("Usuário", placeholder="Digite seu usuário")
            senha = st.text_input("Senha", type="password", placeholder="Digite sua senha")
            st.markdown(
                '<div class="happ-remember-row">'
                '<span>☐ Lembrar de mim</span>'
                '<span class="happ-forgot">Esqueceu a senha?</span>'
                '</div>',
                unsafe_allow_html=True,
            )
            entrar = st.form_submit_button("Entrar →", type="primary", use_container_width=True)
        if entrar:
            auth.garantir_bootstrap()
            usuario = usuario.strip()
            dados = auth.autenticar(usuario, senha)
            if not dados:
                st.error("Usuário ou senha inválidos.")
            elif not auth.usuario_esta_ativo(dados):
                st.error("Este usuário está inativo. Fale com o administrador do seu escritório.")
            else:
                token = auth.criar_sessao(usuario)
                # limpa qualquer resto de sessao anterior (outro usuario/
                # escritorio) ANTES de popular a nova - ver
                # _resetar_estado_nao_login pro bug que isso evita.
                _resetar_estado_nao_login()
                _popular_sessao(dados, token=token)
                _agendar_cookie_sessao("set", token)
                st.rerun()

    with col_dir:
        st.markdown(
            """
            <div class="happ-side-tagline">
                <p>Conectando<br/>pessoas, processos<br/>e resultados.</p>
                <div class="happ-underline"></div>
            </div>
            """,
            unsafe_allow_html=True,
        )

    st.markdown(
        """
        <div class="happ-features">
            <div class="happ-feature"><span class="ficon">⚡</span>
                <div><div class="ftitle">Automatize</div><div class="fdesc">tarefas repetitivas</div></div></div>
            <div class="happ-feature"><span class="ficon">🚀</span>
                <div><div class="ftitle">Ganhe tempo</div><div class="fdesc">para o que importa</div></div></div>
            <div class="happ-feature"><span class="ficon">📊</span>
                <div><div class="ftitle">Tenha mais</div><div class="fdesc">produtividade</div></div></div>
        </div>
        <div class="happ-footer">HUB APP&nbsp;&nbsp;|&nbsp;&nbsp;A CONTABILIDADE MAIS INTELIGENTE</div>
        """,
        unsafe_allow_html=True,
    )


def _tela_trocar_senha(obrigatoria: bool) -> None:
    """Tela de troca de senha com a mesma identidade visual do login
    (_estilo_hub_app) - aparece tanto no primeiro acesso (senha
    temporaria, obrigatoria=True) quanto acionada pelo botao "Trocar
    senha" da sidebar (obrigatoria=False, com opcao de cancelar)."""
    _estilo_hub_app()
    subtitulo = (
        "Por segurança, defina uma nova senha antes de continuar "
        "(esta conta ainda está com a senha padrão/temporária)."
        if obrigatoria else
        "Atualize a senha da sua conta."
    )
    st.markdown(
        f"""
        <div class="happ-header">
            <div class="happ-logo-row">
                <div class="happ-logo-icon">🧩</div>
                <div class="happ-logo-text">Hub <span class="happ-app">APP</span></div>
            </div>
            <div class="happ-welcome">🔑 Trocar senha</div>
            <div class="happ-subtitle">{subtitulo}</div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    _esq, meio, _dir = st.columns([1, 1.3, 1])
    with meio:
        with st.form("trocar_senha_form"):
            senha_atual = st.text_input("Senha atual", type="password", placeholder="Digite sua senha atual")
            nova = st.text_input("Nova senha (mín. 6 caracteres)", type="password", placeholder="Digite a nova senha")
            confirmar = st.text_input("Confirmar nova senha", type="password", placeholder="Repita a nova senha")
            enviar = st.form_submit_button("Salvar nova senha →", type="primary", use_container_width=True)
        if enviar:
            usuario = st.session_state["usuario_logado"]
            if not auth.autenticar(usuario, senha_atual):
                st.error("Senha atual incorreta.")
            elif len(nova) < 6:
                st.error("A nova senha precisa ter pelo menos 6 caracteres.")
            elif nova != confirmar:
                st.error("As senhas digitadas não coincidem.")
            else:
                auth.redefinir_senha(usuario, nova, forcar_troca=False)
                st.session_state["deve_trocar_senha"] = False
                st.session_state["mostrar_trocar_senha"] = False
                st.success("Senha atualizada.")
                st.rerun()
        if not obrigatoria and st.button("Cancelar", use_container_width=True):
            st.session_state["mostrar_trocar_senha"] = False
            st.rerun()


def _flash(chave: str, texto: Optional[str] = None) -> None:
    """Guarda uma mensagem de sucesso para sobreviver a um st.rerun()
    (sem isso, a mensagem e substituida pela proxima renderizacao antes do
    usuario conseguir ve-la). Chame com `texto` para guardar, e sem `texto`
    no topo do bloco seguinte para exibir e limpar a mensagem pendente."""
    if texto is not None:
        st.session_state[chave] = texto
        return
    msg = st.session_state.pop(chave, None)
    if msg:
        st.success(msg)


def _tela_gerenciar_escritorios() -> None:
    """Tela dedicada (só super_admin_global) para gerenciar os escritórios
    cadastrados: criar, renomear, ver quantos usuários cada um tem, e
    excluir. Fica separada do painel de usuários porque mexe no tenant em
    si, não em quem tem acesso a ele."""
    st.title("🌐 Gerenciar Escritórios")
    if st.button("← Início"):
        st.session_state["tela"] = "home"
        st.rerun()
    st.caption("Cada escritório é isolado: usuários de um não enxergam nem gerenciam os de outro.")
    st.divider()

    escritorios = auth.carregar_escritorios()
    usuarios = auth.carregar_usuarios()

    st.subheader("Criar novo escritório")
    _flash("flash_novo_escritorio")
    with st.form("novo_escritorio_form", clear_on_submit=True):
        nome_escritorio = st.text_input("Nome do escritório")
        st.caption("Primeiro usuário desse escritório (será admin_escritorio dele):")
        adm_usuario = st.text_input("Usuário do admin do escritório")
        adm_nome = st.text_input("Nome completo do admin (opcional)")
        adm_senha = st.text_input("Senha inicial (mín. 6 caracteres)", type="password")
        criar = st.form_submit_button("Criar escritório", type="primary")
    if criar:
        if not nome_escritorio.strip() or not adm_usuario.strip():
            st.error("Informe o nome do escritório e o usuário do admin.")
        elif len(adm_senha) < 6:
            st.error("A senha precisa ter pelo menos 6 caracteres.")
        elif adm_usuario.strip() in usuarios:
            st.error("Esse nome de usuário já existe — escolha outro.")
        else:
            novo_eid = auth.criar_escritorio(nome_escritorio.strip())
            auth.criar_ou_atualizar_usuario(
                adm_usuario.strip(), adm_senha, escritorio_id=novo_eid,
                papel=auth.PAPEL_ADMIN_ESCRITORIO, nome=adm_nome.strip(), deve_trocar_senha=True,
            )
            _flash("flash_novo_escritorio",
                   f"✅ Escritório '{nome_escritorio.strip()}' criado com sucesso (id: {novo_eid}), "
                   f"com admin '{adm_usuario.strip()}'.")
            st.rerun()

    st.divider()
    st.subheader(f"Escritórios cadastrados ({len(escritorios)})")
    _flash("flash_escritorio_renomeado")
    _flash("flash_escritorio_removido")

    if not escritorios:
        st.info("Nenhum escritório cadastrado ainda.")
        return

    for eid, edados in sorted(escritorios.items(), key=lambda kv: kv[1].get("nome", kv[0]).lower()):
        usuarios_do = auth.usuarios_do_escritorio(usuarios, eid)
        # mantem o expander aberto entre reruns enquanto os checkboxes de
        # exclusao estiverem marcados - sem isso, cada clique dentro dele
        # fecha o painel de novo (expander sem `expanded` volta a False a
        # cada rerun) e o usuario precisa reabrir a cada passo.
        aberto = bool(
            st.session_state.get(f"forcar_excluir_{eid}") or st.session_state.get(f"confirmar_excluir_{eid}")
        )
        with st.expander(f"🏢 {edados.get('nome', eid)}  ·  {len(usuarios_do)} usuário(s)", expanded=aberto):
            st.caption(f"id: `{eid}` · criado em {edados.get('criado_em', '—')}")

            with st.form(f"renomear_{eid}_form"):
                novo_nome = st.text_input("Nome", value=edados.get("nome", eid), key=f"nome_input_{eid}")
                salvar_nome = st.form_submit_button("Salvar nome")
            if salvar_nome:
                if not novo_nome.strip():
                    st.error("O nome não pode ficar em branco.")
                else:
                    auth.renomear_escritorio(eid, novo_nome.strip())
                    if st.session_state.get("escritorio_id") == eid:
                        st.session_state["escritorio_nome"] = novo_nome.strip()
                    _flash("flash_escritorio_renomeado", f"✅ Escritório renomeado para '{novo_nome.strip()}'.")
                    st.rerun()

            st.markdown("**Aplicativos permitidos**")
            st.caption("Nenhum marcado = sem restrição (vê todos). Isso é o teto: cada usuário do "
                       "escritório pode ainda ser restrito mais em Gerenciar Usuários, mas nunca além disso.")
            _ids_apps = [a["id"] for a in _APPS_HOME]
            _apps_atuais_escritorio = [a for a in edados.get("apps_permitidos") or [] if a in _ids_apps]
            _apps_escolhidos_escritorio = st.multiselect(
                "Apps", options=_ids_apps,
                format_func=lambda aid: next((a["titulo"] for a in _APPS_HOME if a["id"] == aid), aid),
                default=_apps_atuais_escritorio, key=f"apps_escritorio_{eid}", label_visibility="collapsed",
            )
            if st.button("Salvar apps permitidos", key=f"salvar_apps_escritorio_{eid}"):
                auth.definir_apps_escritorio(eid, _apps_escolhidos_escritorio)
                _flash("flash_escritorio_renomeado", "✅ Apps permitidos do escritório atualizados.")
                st.rerun()

            st.markdown("**Licença do script local (ISS Net Online attended)**")
            st.caption("Se desmarcar, o script `attended_worker/issnet_attended.py` recusa a rodar "
                       "pra este escritório na próxima vez que for usado.")
            _liberado_atual = edados.get("issnet_attended_liberado", True)
            _liberado_novo = st.checkbox(
                "Liberado pra usar o script attended", value=_liberado_atual, key=f"licenca_attended_{eid}",
            )
            if _liberado_novo != _liberado_atual:
                auth.definir_issnet_attended_liberado(eid, _liberado_novo)
                _flash(
                    "flash_escritorio_renomeado",
                    "✅ Licença do script attended " + ("liberada." if _liberado_novo else "bloqueada."),
                )
                st.rerun()

            st.markdown("**Usuários deste escritório**")
            if usuarios_do:
                for uname, udados in usuarios_do.items():
                    st.write(f"- **{uname}** — {udados.get('nome', uname)} · {udados.get('papel')}")
            else:
                st.caption("Nenhum usuário cadastrado neste escritório.")

            st.markdown("**Excluir escritório**")
            forcar = False
            if usuarios_do:
                st.warning(f"Este escritório tem {len(usuarios_do)} usuário(s) cadastrado(s).")
                forcar = st.checkbox(
                    "Excluir também todos os usuários deste escritório", key=f"forcar_excluir_{eid}"
                )
            confirmar = st.checkbox(f"Confirmo que quero excluir \"{edados.get('nome', eid)}\"", key=f"confirmar_excluir_{eid}")
            if st.button("Excluir escritório", key=f"btn_excluir_{eid}", disabled=not confirmar):
                ok, msg = auth.remover_escritorio(eid, forcar=forcar)
                if ok:
                    _flash("flash_escritorio_removido", f"✅ {msg}")
                    st.rerun()
                else:
                    st.error(msg)


def _tela_gerenciar_usuarios() -> None:
    """Tela dedicada para gerenciar os usuários cadastrados: cadastrar,
    inativar/ativar (bloqueia o login sem apagar o cadastro) e remover
    definitivamente. super_admin_global enxerga e gerencia usuários de
    TODOS os escritórios (escolhendo qual ver); admin_escritorio só
    enxerga/gerencia o próprio (isolamento entre escritórios-cliente)."""
    st.title("👥 Gerenciar Usuários")
    if st.button("← Início"):
        st.session_state["tela"] = "home"
        st.rerun()
    st.divider()

    eh_global = st.session_state.get("papel_usuario") == auth.PAPEL_SUPER_GLOBAL
    usuarios_todos = auth.carregar_usuarios()
    escritorios = auth.carregar_escritorios()

    if eh_global:
        if not escritorios:
            st.info("Nenhum escritório cadastrado ainda — crie um em "
                     "\"🌐 Gerenciar Escritórios\" antes de cadastrar usuários.")
            return
        escritorio_visto = st.selectbox(
            "Escritório",
            list(escritorios.keys()),
            format_func=lambda eid: escritorios.get(eid, {}).get("nome", eid),
            index=list(escritorios.keys()).index(st.session_state.get("escritorio_id"))
            if st.session_state.get("escritorio_id") in escritorios else 0,
        )
        papeis_disponiveis = [auth.PAPEL_USUARIO, auth.PAPEL_ADMIN_ESCRITORIO, auth.PAPEL_SUPER_GLOBAL]
    else:
        escritorio_visto = st.session_state.get("escritorio_id")
        papeis_disponiveis = [auth.PAPEL_USUARIO, auth.PAPEL_ADMIN_ESCRITORIO]
        st.caption(f"Escritório: **{st.session_state.get('escritorio_nome')}**")

    usuarios_escopo = auth.usuarios_do_escritorio(usuarios_todos, escritorio_visto)

    st.subheader("Cadastrar usuário / redefinir senha")
    _flash("flash_usuario_salvo")
    with st.form("gerenciar_usuario_form", clear_on_submit=True):
        alvo = st.text_input("Usuário (novo ou existente)")
        nome_completo = st.text_input("Nome completo (opcional)")
        papel = st.selectbox("Papel", papeis_disponiveis)
        nova_senha = st.text_input("Senha (mín. 6 caracteres)", type="password")
        salvar = st.form_submit_button("Salvar", type="primary")
    if salvar:
        alvo = alvo.strip()
        ja_existe = alvo in usuarios_todos
        if not alvo:
            st.error("Informe o nome de usuário.")
        elif len(nova_senha) < 6:
            st.error("A senha precisa ter pelo menos 6 caracteres.")
        elif ja_existe and not eh_global and usuarios_todos[alvo].get("escritorio_id") != escritorio_visto:
            st.error("Esse usuário pertence a outro escritório — você só gerencia o seu.")
        else:
            auth.criar_ou_atualizar_usuario(
                alvo, nova_senha, escritorio_id=escritorio_visto, papel=papel,
                nome=nome_completo.strip(), deve_trocar_senha=True,
            )
            _flash("flash_usuario_salvo",
                   f"✅ Usuário '{alvo}' salvo com sucesso — vai precisar trocar a senha no próximo login.")
            st.rerun()

    st.divider()
    st.subheader(f"Usuários cadastrados ({len(usuarios_escopo)})")
    _flash("flash_usuario_status")
    _flash("flash_usuario_removido")

    if not usuarios_escopo:
        st.info("Nenhum usuário cadastrado neste escritório ainda.")
        return

    for uname, dados in sorted(usuarios_escopo.items()):
        ativo = auth.usuario_esta_ativo(dados)
        voce = uname == st.session_state.get("usuario_logado")
        status_txt = "🟢 Ativo" if ativo else "🔴 Inativo"
        # mantem aberto entre reruns enquanto houver uma acao pendente
        # marcada neste usuario (mesmo motivo do fix nos escritorios).
        aberto = bool(
            st.session_state.get(f"confirmar_remover_{uname}")
            or st.session_state.get(f"aberto_{uname}")
        )
        rotulo = f"{dados.get('nome', uname)} (`{uname}`) · {dados.get('papel')} · {status_txt}"
        with st.expander(rotulo, expanded=aberto):
            st.caption(f"Papel: {dados.get('papel')}" + (" · (você)" if voce else ""))

            c1, c2 = st.columns(2)
            with c1:
                if ativo:
                    if voce:
                        st.caption("Você não pode inativar a própria conta.")
                    elif st.button("🔴 Inativar", key=f"inativar_{uname}"):
                        ok, msg = auth.definir_status_usuario(uname, ativo=False)
                        if ok:
                            _flash("flash_usuario_status", f"✅ {msg}")
                            st.rerun()
                        else:
                            st.error(msg)
                else:
                    if st.button("🟢 Ativar", key=f"ativar_{uname}", type="primary"):
                        ok, msg = auth.definir_status_usuario(uname, ativo=True)
                        if ok:
                            _flash("flash_usuario_status", f"✅ {msg}")
                            st.rerun()
                        else:
                            st.error(msg)
            with c2:
                if voce:
                    st.caption("Você não pode remover a própria conta.")
                else:
                    confirmar = st.checkbox("Confirmar remoção definitiva", key=f"confirmar_remover_{uname}")
                    if st.button("🗑️ Remover definitivamente", key=f"remover_{uname}", disabled=not confirmar):
                        if auth.remover_usuario(uname):
                            _flash("flash_usuario_removido", f"✅ Usuário '{uname}' removido definitivamente.")
                            st.rerun()
                        else:
                            st.error("Não foi possível remover: precisa sobrar pelo menos 1 admin para este escritório.")

            st.divider()
            st.markdown("**Aplicativos permitidos**")
            _ids_apps_escritorio_do_user = (
                escritorios.get(dados.get("escritorio_id"), {}).get("apps_permitidos")
                or [a["id"] for a in _APPS_HOME]
            )
            st.caption("Nenhum marcado = vê tudo que o escritório permite. As opções aqui já respeitam "
                       "o teto definido em Gerenciar Escritórios.")
            _apps_atuais_user = [a for a in dados.get("apps_permitidos") or [] if a in _ids_apps_escritorio_do_user]
            _apps_escolhidos_user = st.multiselect(
                "Apps", options=_ids_apps_escritorio_do_user,
                format_func=lambda aid: next((a["titulo"] for a in _APPS_HOME if a["id"] == aid), aid),
                default=_apps_atuais_user, key=f"apps_usuario_{uname}", label_visibility="collapsed",
            )
            if st.button("Salvar apps permitidos", key=f"salvar_apps_usuario_{uname}"):
                auth.definir_apps_usuario(uname, _apps_escolhidos_user)
                _flash("flash_usuario_status", "✅ Apps permitidos do usuário atualizados.")
                st.rerun()

            st.divider()
            st.markdown("**Redefinir senha**")
            with st.form(f"redefinir_senha_{uname}_form", clear_on_submit=True):
                nova_senha_user = st.text_input(
                    "Nova senha (mín. 6 caracteres)", type="password", key=f"nova_senha_{uname}"
                )
                redefinir = st.form_submit_button("🔑 Redefinir senha")
            if redefinir:
                if len(nova_senha_user) < 6:
                    st.error("A senha precisa ter pelo menos 6 caracteres.")
                    st.session_state[f"aberto_{uname}"] = True
                else:
                    auth.redefinir_senha(uname, nova_senha_user, forcar_troca=True)
                    st.session_state[f"aberto_{uname}"] = False
                    _flash("flash_usuario_status",
                           f"✅ Senha de '{uname}' redefinida — vai precisar trocar no próximo login.")
                    st.rerun()


def _tela_historico() -> None:
    """Tela de historico dos lancamentos (conciliacoes) rodados por cada
    usuario - qualquer usuario logado enxerga (inclusive usuario comum),
    sempre restrito ao PROPRIO escritorio (historico.listar ja filtra por
    escritorio_id); so super_admin_global tem o seletor pra escolher
    QUALQUER escritorio, por ser quem administra o hub inteiro."""
    st.title("📜 Histórico de Lançamentos")
    if st.button("← Início"):
        st.session_state["tela"] = "home"
        st.rerun()
    st.divider()

    eh_global = st.session_state.get("papel_usuario") == auth.PAPEL_SUPER_GLOBAL
    if eh_global:
        escritorios = auth.carregar_escritorios()
        if not escritorios:
            st.info("Nenhum escritório cadastrado ainda.")
            return
        escritorio_visto = st.selectbox(
            "Escritório",
            list(escritorios.keys()),
            format_func=lambda eid: escritorios.get(eid, {}).get("nome", eid),
            index=list(escritorios.keys()).index(st.session_state.get("escritorio_id"))
            if st.session_state.get("escritorio_id") in escritorios else 0,
        )
    else:
        escritorio_visto = st.session_state.get("escritorio_id")
        st.caption(f"Escritório: **{st.session_state.get('escritorio_nome')}**")

    registros = historico.listar(escritorio_visto)
    if not registros:
        st.info("Nenhum lançamento registrado ainda neste escritório.")
        return

    st.caption(f"{len(registros)} lançamento(s) mais recente(s) primeiro:")
    for r in registros:
        dt = r["criado_em"]
        empresa = r.get("empresa_nome") or r["empresa_codigo"]
        st.write(
            f"👤 **{r['usuario']}** — Lançamento {empresa} "
            f"dia {dt.strftime('%d/%m/%Y')} hora {dt.strftime('%H:%M')}"
        )


def _tela_gerenciar_clientes() -> None:
    """Tela dedicada para gerenciar os clientes (empresas) cadastrados:
    cadastrar, editar nome/código e remover. Disponível pra qualquer
    usuário logado (é só referência de negócio, não dado sensível) -
    super_admin_global escolhe qual escritório ver; os demais só enxergam
    o próprio. Remoção fica restrita a admin_escritorio/super_admin_global
    pra evitar que alguém apague por engano um código usado por outros."""
    st.title("🏢 Gerenciar Clientes")
    if st.button("← Início"):
        st.session_state["tela"] = "home"
        st.rerun()
    st.divider()

    eh_global = st.session_state.get("papel_usuario") == auth.PAPEL_SUPER_GLOBAL
    pode_remover = st.session_state.get("papel_usuario") in (auth.PAPEL_SUPER_GLOBAL, auth.PAPEL_ADMIN_ESCRITORIO)

    if eh_global:
        escritorios = auth.carregar_escritorios()
        if not escritorios:
            st.info("Nenhum escritório cadastrado ainda.")
            return
        escritorio_visto = st.selectbox(
            "Escritório",
            list(escritorios.keys()),
            format_func=lambda eid: escritorios.get(eid, {}).get("nome", eid),
            index=list(escritorios.keys()).index(st.session_state.get("escritorio_id"))
            if st.session_state.get("escritorio_id") in escritorios else 0,
        )
    else:
        escritorio_visto = st.session_state.get("escritorio_id")
        st.caption(f"Escritório: **{st.session_state.get('escritorio_nome')}**")

    st.subheader("Cadastrar cliente")
    _flash("flash_gerenciar_cliente")
    with st.form("gerenciar_cliente_form", clear_on_submit=True):
        novo_nome = st.text_input("Nome do cliente")
        novo_codigo = st.text_input("Código da empresa no Domínio")
        salvar = st.form_submit_button("Cadastrar", type="primary")
    if salvar:
        if not novo_nome.strip() or not novo_codigo.strip():
            st.error("Informe nome e código da empresa.")
        else:
            ok, msg = clientes.criar_cliente(escritorio_visto, novo_nome.strip(), novo_codigo.strip())
            if ok:
                _flash("flash_gerenciar_cliente", f"✅ {msg}")
                st.rerun()
            else:
                st.error(msg)

    st.divider()
    lista = clientes.listar_clientes(escritorio_visto)
    st.subheader(f"Clientes cadastrados ({len(lista)})")

    if not lista:
        st.info("Nenhum cliente cadastrado neste escritório ainda.")
        return

    for c in lista:
        aberto = bool(st.session_state.get(f"aberto_cliente_{c['id']}"))
        with st.expander(f"{c['nome']} — {c['codigo_dominio']}", expanded=aberto):
            with st.form(f"editar_cliente_{c['id']}_form"):
                nome_edit = st.text_input("Nome", value=c["nome"], key=f"nome_cliente_{c['id']}")
                codigo_edit = st.text_input(
                    "Código da empresa no Domínio", value=c["codigo_dominio"], key=f"codigo_cliente_{c['id']}"
                )
                salvar_edit = st.form_submit_button("💾 Salvar alterações")
            if salvar_edit:
                if not nome_edit.strip() or not codigo_edit.strip():
                    st.error("Nome e código não podem ficar em branco.")
                    st.session_state[f"aberto_cliente_{c['id']}"] = True
                else:
                    ok, msg = clientes.atualizar_cliente(
                        c["id"], escritorio_visto, nome_edit.strip(), codigo_edit.strip()
                    )
                    if ok:
                        _flash("flash_gerenciar_cliente", f"✅ {msg}")
                        st.session_state[f"aberto_cliente_{c['id']}"] = False
                        st.rerun()
                    else:
                        st.error(msg)
                        st.session_state[f"aberto_cliente_{c['id']}"] = True

            if pode_remover:
                confirmar = st.checkbox(
                    "Confirmar remoção definitiva", key=f"confirmar_remover_cliente_{c['id']}"
                )
                if confirmar:
                    st.session_state[f"aberto_cliente_{c['id']}"] = True
                if st.button(
                    "🗑️ Remover definitivamente", key=f"remover_cliente_{c['id']}", disabled=not confirmar,
                ):
                    clientes.remover_cliente(c["id"], escritorio_visto)
                    st.session_state[f"aberto_cliente_{c['id']}"] = False
                    _flash("flash_gerenciar_cliente", f"✅ Cliente '{c['nome']}' removido.")
                    st.rerun()


_CARACTERES_INVALIDOS_PASTA = str.maketrans({c: "_" for c in '/\\:*?"<>|'})


def _contar_arquivos_zip(conteudo: bytes) -> int:
    """Quantos arquivos tem dentro de um .zip de XMLs — usado só pra mostrar
    'quantidade de notas baixadas' na tela, sem abrir/extrair nada."""
    try:
        with zipfile.ZipFile(io.BytesIO(conteudo)) as zf:
            return len([n for n in zf.namelist() if not n.endswith("/")])
    except zipfile.BadZipFile:
        return 0


def _nome_pasta_empresa(empresa: dict) -> str:
    """'<código> - <razão social>' (como pede a especificação do issnet) —
    cai só no código quando não há razão social (ex.: issweb não captura)."""
    razao = (empresa.get("razao_social") or "").strip()
    nome = f"{empresa['codigo']} - {razao}" if razao else str(empresa["codigo"])
    return nome.translate(_CARACTERES_INVALIDOS_PASTA)


def _planilha_modelo(modulo_info: dict) -> bytes:
    """Planilha em branco (só cabeçalho + 1 linha de exemplo) com as colunas
    esperadas do módulo selecionado, pra baixar e preencher com os dados
    reais — evita ter que descobrir o nome exato das colunas na mão."""
    from openpyxl import Workbook

    colunas = [c.replace(" (opcional)", "") for c in modulo_info["colunas_planilha"]]
    wb = Workbook()
    ws = wb.active
    ws.title = "Empresas"
    ws.append(colunas)
    linha_exemplo = {
        "Código da Empresa": "68", "CNPJ/CPF": "00.000.000/0001-00",
        "Razão Social": "Empresa Exemplo LTDA", "Município": modulo_info["municipio_alvo"].split(" / ")[0],
        "CNPJ": "00.000.000/0001-00", "Inscrição Estadual": "10.123.456-7",
    }
    ws.append([linha_exemplo.get(c, "") for c in colunas])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _gerar_relatorio_geral(execucao: dict, empresas_exec: list[dict]) -> bytes:
    """Relatorio_Geral_Processamento.xlsx: uma linha por empresa/obrigação
    processada na execução, com status, movimento, quantidade de notas
    baixadas e erro (quando houver) — visão consolidada pra quem só quer
    conferir o lote sem abrir cada PDF/XML."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Processamento"
    ws.append(["Código", "Razão Social", "CNPJ/CPF", "Obrigação", "Status", "Movimento", "Notas (XML)", "Erro"])
    for empresa in empresas_exec:
        qtd_xml = _contar_arquivos_zip(bytes(empresa["xml_zip"])) if empresa.get("xml_zip") else 0
        ws.append([
            empresa["codigo"], empresa.get("razao_social") or "", empresa["cnpj_cpf"],
            empresa["obrigacao"], empresa["status"], empresa.get("movimento") or "",
            qtd_xml, empresa.get("erro") or "",
        ])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _montar_zip_execucao(execucao: dict, empresas_exec: list[dict]) -> bytes:
    """Monta um .zip só com os arquivos já concluídos da execução, uma
    pasta por empresa ('<código> - <razão social>') e dentro dela uma
    subpasta pela competência no formato MMAAAA (ex.: '10/082026/'). PDF e
    XML (quando houver) do mesmo jeito que a tela oferece pra baixar
    individualmente, mais um Relatorio_Geral_Processamento.xlsx na raiz do
    zip com o resumo de todas as empresas da execução (concluídas ou não)."""
    competencia = execucao.get("competencia") or ""
    pasta_competencia = competencia.replace("/", "") if competencia else "sem-competencia"

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        for empresa in empresas_exec:
            if empresa["status"] != rpa_core.STATUS_CONCLUIDO:
                continue
            pasta = f"{_nome_pasta_empresa(empresa)}/{pasta_competencia}"
            if empresa.get("pdf"):
                zf.writestr(f"{pasta}/{empresa['pdf_nome']}", bytes(empresa["pdf"]))
            if empresa.get("xml_zip"):
                zf.writestr(f"{pasta}/{empresa['xml_zip_nome']}", bytes(empresa["xml_zip"]))
        zf.writestr("Relatorio_Geral_Processamento.xlsx", _gerar_relatorio_geral(execucao, empresas_exec))
    return buffer.getvalue()


def _interpolar_cor(c1: tuple[int, int, int], c2: tuple[int, int, int], t: float) -> tuple[int, int, int]:
    return tuple(round(c1[i] + (c2[i] - c1[i]) * t) for i in range(3))


def _barra_retencao_arquivos(concluido_em) -> None:
    """Avisa quantos dias faltam até o PDF/XML dessa execução serem
    apagados por retenção (rpa_core.RETENCAO_ARQUIVOS_DIAS dias após
    concluído - ver rpa_core.limpar_arquivos_vencidos, quem apaga de
    fato). Barra que começa verde e vai virando vermelha conforme o
    prazo se aproxima, pra avisar antes do "Baixar tudo" sumir."""
    if not concluido_em:
        return
    total_dias = rpa_core.RETENCAO_ARQUIVOS_DIAS
    agora = datetime.now(timezone.utc)
    if concluido_em.tzinfo is None:
        concluido_em = concluido_em.replace(tzinfo=timezone.utc)
    dias_passados = (agora - concluido_em).total_seconds() / 86400
    fracao = max(0.0, min(1.0, dias_passados / total_dias))
    dias_restantes = max(0, total_dias - int(dias_passados))

    verde, amarelo, vermelho = (34, 197, 94), (234, 179, 8), (239, 68, 68)
    if fracao <= 0.5:
        cor = _interpolar_cor(verde, amarelo, fracao / 0.5)
    else:
        cor = _interpolar_cor(amarelo, vermelho, (fracao - 0.5) / 0.5)

    st.markdown(
        f"""
        <div style="margin:4px 0 10px 0;">
          <div style="font-size:0.8em;color:#666;margin-bottom:2px;">
            🗂️ Arquivos serão apagados em {dias_restantes} dia(s) (retenção de {total_dias} dias)
          </div>
          <div style="background:#e5e7eb;border-radius:6px;height:8px;width:100%;overflow:hidden;">
            <div style="background:rgb{cor};height:100%;width:{fracao * 100:.1f}%;"></div>
          </div>
        </div>
        """,
        unsafe_allow_html=True,
    )


def _tela_rpa_manual(modulo_id: str, modulo_info: dict, escritorio_id: str, usuario: str) -> None:
    """Fluxo de apoio pra rotinas ainda não automatizadas (ver
    modulo_info['automatizado'] em rpa/registry.py): não roda nada
    sozinha - só reaproveita rpa_execucoes/rpa_empresas pra guardar a
    lista de empresas da planilha do mês, com link do portal pra abrir
    numa aba separada, e a pessoa marca cada empresa como feita à mão."""
    if modulo_id == "issnet_rest_dms":
        from rpa.issnet.urls import PORTAL_URLS

        st.markdown("**Script de automação assistida (roda no seu PC)**")
        st.caption(
            "O Cloudflare do portal bloqueia navegador controlado por automação de servidor — "
            "esse programa mecaniza os cliques DEPOIS que você loga manualmente (certificado). "
            "Baixe, extraia e dê duplo clique no .exe: abre uma janela onde você escolhe o "
            "município, faz login no portal e informa a planilha/pasta — sem precisar de terminal."
        )
        _pasta_attended = Path(__file__).parent / "attended_worker"
        _zip_attended = io.BytesIO()
        with zipfile.ZipFile(_zip_attended, "w", zipfile.ZIP_DEFLATED) as zf:
            for _nome_arquivo in ["dist/issnet_attended.exe", "LEIA-ME.txt"]:
                _caminho = _pasta_attended / _nome_arquivo
                if _caminho.exists():
                    zf.write(_caminho, arcname=Path(_nome_arquivo).name)
        st.download_button(
            "📥 Baixar script de automação (.zip)",
            _zip_attended.getvalue(),
            file_name="issnet_attended.zip",
            type="primary",
        )
        st.caption(
            "O .zip traz um .exe pronto — não precisa instalar Python. O Windows pode avisar "
            "'aplicativo desconhecido' (SmartScreen) por não ter certificado de editor; clique em "
            "'Mais informações' → 'Executar assim mesmo'."
        )

        st.divider()
        st.markdown("**Abrir o portal:**")
        cols_link = st.columns(len(PORTAL_URLS))
        for col, (municipio, url) in zip(cols_link, PORTAL_URLS.items()):
            col.link_button(f"🔗 {municipio.title()}", url, use_container_width=True)

        st.divider()
        st.markdown("**Pasta onde os arquivos são salvos (script attended local)**")
        _escritorio_atual = auth.carregar_escritorios().get(escritorio_id, {})
        with st.form(f"pasta_raiz_{modulo_id}_form"):
            _pasta_raiz = st.text_input(
                "Pasta raiz no PC que roda o script",
                value=_escritorio_atual.get("pasta_raiz_local", ""),
                placeholder=r"C:\Prefeituras",
                help="Dentro dela, cada empresa vira uma pasta com o código (da planilha) e, "
                     "dentro dessa, uma subpasta com a competência (ex.: C:\\Prefeituras\\320\\082026\\).",
            )
            _salvar_pasta = st.form_submit_button("💾 Salvar pasta")
        if _salvar_pasta:
            auth.definir_pasta_raiz_local(escritorio_id, _pasta_raiz)
            _flash("flash_rpa_manual", "✅ Pasta raiz salva.")
            st.rerun()

    st.divider()
    st.subheader("Empresas do mês")
    _periodo_padrao = rpa_registry.preparar_periodo(modulo_id)
    _data_competencia = st.date_input(
        "Competência", value=date(_periodo_padrao["ano"], _periodo_padrao["mes"], 1),
        key=f"competencia_manual_{modulo_id}", format="DD/MM/YYYY",
    )
    competencia_escolhida = f"{_data_competencia.month:02d}/{_data_competencia.year}"

    up_planilha = st.file_uploader(
        "Planilha de empresas (.xlsx)", type=["xlsx"], key=f"up_planilha_manual_{modulo_id}",
    )
    if up_planilha:
        conteudo = up_planilha.getvalue()
        try:
            empresas = rpa_registry.ler_empresas(modulo_id, conteudo)
        except Exception as exc:
            st.error(str(exc))
        else:
            st.success(f"{len(empresas)} empresa(s) encontradas na planilha.")
            if st.button("📋 Criar lista de controle", type="primary"):
                execucao_id = rpa_core.criar_execucao(
                    escritorio_id, modulo_id, conteudo, up_planilha.name, usuario, empresas,
                    competencia_escolhida,
                )
                _flash(
                    "flash_rpa_manual",
                    f"✅ Lista #{execucao_id} criada — {len(empresas)} empresa(s), competência {competencia_escolhida}.",
                )
                st.rerun()

    st.divider()
    _flash("flash_rpa_manual")
    st.subheader("Listas de controle")

    # BOTAO TEMPORARIO DE MANUTENCAO - pedido explicito do usuario pra
    # apagar TODAS as listas (de qualquer escritorio) e reiniciar a
    # numeracao do zero, pra testar se o isolamento por escritorio
    # (listar_execucoes ja filtra por escritorio_id) aparece certo numa
    # lista vazia. So super_admin_global ve, e exige dois cliques
    # (checkbox de confirmacao + botao) pra nao apagar sem querer.
    # REMOVER essa caixa inteira depois de usar uma vez - ver
    # rpa_core.limpar_execucoes_modulo.
    if st.session_state.get("papel_usuario") == auth.PAPEL_SUPER_GLOBAL:
        with st.expander("🛠️ Manutenção (temporário) — apagar TODAS as listas deste módulo"):
            st.warning(
                "Apaga TODAS as execuções/listas do módulo (de QUALQUER escritório, não só o "
                "selecionado) e os PDFs/XMLs junto — irreversível. Só pra reiniciar a numeração "
                "em teste."
            )
            _confirma_limpar = st.checkbox(
                "Confirmo que quero apagar TUDO deste módulo, de todos os escritórios",
                key=f"confirma_limpar_tudo_{modulo_id}",
            )
            if st.button("🗑️ Apagar todas as listas agora", disabled=not _confirma_limpar, key=f"btn_limpar_tudo_{modulo_id}"):
                qtd = rpa_core.limpar_execucoes_modulo(modulo_id)
                _flash("flash_rpa_manual", f"✅ {qtd} lista(s) apagada(s) — numeração reiniciada.")
                st.rerun()

    rpa_core.limpar_arquivos_vencidos()
    execucoes = rpa_core.listar_execucoes(escritorio_id, modulo_id)
    if not execucoes:
        st.info("Nenhuma lista criada ainda para esta rotina.")
        return

    _emoji_status_manual = {
        rpa_core.STATUS_PENDENTE: "⏳ Pendente", rpa_core.STATUS_RODANDO: "🔄 Rodando",
        rpa_core.STATUS_CONCLUIDO: "✅ Feito", rpa_core.STATUS_ERRO: "❌ Erro",
    }
    for execucao in execucoes:
        empresas_exec = rpa_core.listar_empresas(execucao["id"])
        concluidas = sum(1 for e in empresas_exec if e["status"] == rpa_core.STATUS_CONCLUIDO)
        erros = sum(1 for e in empresas_exec if e["status"] == rpa_core.STATUS_ERRO)
        # inclui PENDENTE/RODANDO travado numa execução que já não está
        # mais PENDENTE (processamento interrompido no meio) - essas
        # empresas nunca chegaram a rodar, mas a API do attended só
        # procura execução com status PENDENTE, então ficam invisíveis
        # pra sempre sem reabrir a execução (ver reabrir_execucao_se_incompleta).
        pendencias = sum(1 for e in empresas_exec if e["status"] != rpa_core.STATUS_CONCLUIDO)
        _competencia_exec = f" · competência {execucao['competencia']}" if execucao.get("competencia") else ""
        titulo = (
            f"📋 Lista #{execucao['id']} — {execucao['criado_em']:%d/%m/%Y %H:%M}{_competencia_exec} "
            f"({concluidas}/{len(empresas_exec)} feitas{f', {erros} erro(s)' if erros else ''})"
        )
        with st.expander(titulo):
            st.caption(f"Planilha: {execucao['planilha_nome']} · Enviada por {execucao['criado_por']}")
            if concluidas and execucao.get("concluido_em"):
                _barra_retencao_arquivos(execucao["concluido_em"])
            _cols_acoes_manual = st.columns(2)
            if concluidas:
                _pasta_comp_manual = (execucao.get("competencia") or "sem-competencia").replace("/", "")
                _cols_acoes_manual[0].download_button(
                    "📦 Baixar tudo (.zip)", _montar_zip_execucao(execucao, empresas_exec),
                    file_name=f"{modulo_id} {_pasta_comp_manual}.zip", key=f"zip_manual_{execucao['id']}",
                )
            if pendencias:
                if _cols_acoes_manual[1].button(
                    f"🔁 Reprocessar {pendencias} empresa(s) pendente(s) (via Sincronizar com o Hub)",
                    key=f"reprocessar_manual_{execucao['id']}",
                ):
                    rpa_core.reprocessar_falhas(execucao["id"])
                    rpa_core.reabrir_execucao_se_incompleta(execucao["id"])
                    _flash(
                        "flash_rpa_manual",
                        f"✅ Lista reaberta — rode o script attended de novo (Sincronizar com o "
                        "Hub) que ele processa só quem ainda não terminou.",
                    )
                    st.rerun()
            for empresa in empresas_exec:
                status_atual = empresa["status"]
                feito = status_atual == rpa_core.STATUS_CONCLUIDO
                cols = st.columns([1, 2, 2, 2, 1, 1, 1])
                cols[0].write(empresa["codigo"])
                cols[1].write(empresa["cnpj_cpf"])
                cols[2].write(empresa.get("municipio") or "—")
                cols[3].write(_emoji_status_manual.get(status_atual, status_atual))
                if empresa.get("erro"):
                    cols[3].caption(empresa["erro"][:120])
                if empresa.get("pdf"):
                    cols[4].download_button(
                        "PDF", bytes(empresa["pdf"]), file_name=empresa["pdf_nome"],
                        key=f"pdf_manual_{empresa['id']}",
                    )
                if empresa.get("xml_zip"):
                    cols[5].download_button(
                        "XML", bytes(empresa["xml_zip"]), file_name=empresa["xml_zip_nome"],
                        key=f"xml_manual_{empresa['id']}",
                    )
                if feito:
                    if cols[6].button("↩️", key=f"desfazer_manual_{empresa['id']}", help="Desfazer"):
                        rpa_core.marcar_empresa_status(empresa["id"], rpa_core.STATUS_PENDENTE)
                        st.rerun()
                else:
                    if cols[6].button("✔️", key=f"concluir_manual_{empresa['id']}", help="Marcar como feito"):
                        rpa_core.marcar_empresa_status(empresa["id"], rpa_core.STATUS_CONCLUIDO)
                        st.rerun()


_STATUS_NFGO = {
    rpa_core.STATUS_PENDENTE: "⏳ Aguardando", rpa_core.STATUS_RODANDO: "🔄 Executando",
    rpa_core.STATUS_ERRO: "❌ Erro",
}


def _status_linha_nfgo(linha: dict | None) -> str:
    if linha is None:
        return "—"
    if linha["status"] == rpa_core.STATUS_CONCLUIDO:
        return "➖ Sem movimento" if linha.get("movimento") == "Sem movimento" else "✅ Concluído"
    return _STATUS_NFGO.get(linha["status"], linha["status"])


def _status_empresa_nfgo(linhas: list[dict]) -> str:
    """Status consolidado da empresa (Entrada + Saída): o pior dos dois."""
    status = [_status_linha_nfgo(l) for l in linhas]
    for rotulo in ("❌ Erro", "🔄 Executando", "⏳ Aguardando"):
        if rotulo in status:
            return rotulo
    if all(s == "➖ Sem movimento" for s in status):
        return "➖ Sem movimento"
    return "✅ Concluído"


def _agrupar_nfgo(empresas_exec: list[dict]) -> list[tuple[dict, dict | None, dict | None]]:
    """[(dados da empresa, linha ENTRADA, linha SAIDA)] na ordem da planilha."""
    grupos: dict = {}
    for e in empresas_exec:
        chave = (e["codigo"], e.get("inscricao_estadual") or "")
        grupo = grupos.setdefault(chave, {"empresa": e, "ENTRADA": None, "SAIDA": None})
        grupo[(e.get("obrigacao") or "").upper()] = e
    return [(g["empresa"], g["ENTRADA"], g["SAIDA"]) for g in grupos.values()]


def _qtd_xml_nfgo(linha: dict | None):
    if linha is None or linha["status"] in (rpa_core.STATUS_PENDENTE, rpa_core.STATUS_RODANDO):
        return None
    return linha.get("qtd_xml") if linha.get("qtd_xml") is not None else (0 if linha.get("movimento") == "Sem movimento" else None)


def _linhas_grade_nfgo(execucao: dict, empresas_exec: list[dict]) -> list[dict]:
    grade = []
    for empresa, entrada, saida in _agrupar_nfgo(empresas_exec):
        qtd_e, qtd_s = _qtd_xml_nfgo(entrada), _qtd_xml_nfgo(saida)
        atualizados = [l["atualizado_em"] for l in (entrada, saida) if l and l.get("atualizado_em")]
        observacoes = [
            f"{tipo}: {l.get('observacao') or l.get('erro')}"
            for tipo, l in (("Entrada", entrada), ("Saída", saida))
            if l and (l.get("observacao") or l.get("erro"))
        ]
        grade.append({
            "Código": empresa["codigo"],
            "Empresa": empresa.get("razao_social") or "",
            "CNPJ": empresa["cnpj_cpf"],
            "IE": empresa.get("inscricao_estadual") or "",
            "Competência": execucao.get("competencia") or "",
            "XML Entrada": qtd_e,
            "SEFAZ Entrada": entrada.get("qtd_notas_portal") if entrada else None,
            "XML Saída": qtd_s,
            "SEFAZ Saída": saida.get("qtd_notas_portal") if saida else None,
            "Total": (qtd_e or 0) + (qtd_s or 0) if qtd_e is not None or qtd_s is not None else None,
            "Status": _status_empresa_nfgo([l for l in (entrada, saida) if l]),
            "Data/Hora": max(atualizados).strftime("%d/%m/%Y %H:%M") if atualizados else "",
            "Observação": " | ".join(observacoes),
        })
    # célula vazia em vez de "None" (consulta ainda não feita)
    return [{k: ("" if v is None else v) for k, v in linha.items()} for linha in grade]


def _montar_zip_nfgo(execucao: dict, empresas_exec: list[dict]) -> bytes:
    """Mesma estrutura que o worker grava em disco (ver
    rpa/sefazgo_nfe/arquivos.py pasta_relativa): RPA NF GO/<código -
    empresa>/<MMAAAA>/ENTRADA|SAIDA/ com o ZIP de XMLs e o print da
    consulta, mais a planilha-resumo na raiz."""
    competencia = execucao.get("competencia") or ""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        for e in empresas_exec:
            tipo = (e.get("obrigacao") or "").upper()
            pasta = nfgo_arquivos.pasta_relativa(e["codigo"], e.get("razao_social") or "", competencia, tipo)
            if e.get("xml_zip"):
                zf.writestr(f"{pasta}/{e['xml_zip_nome']}", bytes(e["xml_zip"]))
            if e.get("evidencia_png"):
                zf.writestr(f"{pasta}/{nfgo_arquivos.nome_evidencia(tipo, competencia)}", bytes(e["evidencia_png"]))

        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "RPA NF GO"
        grade = _linhas_grade_nfgo(execucao, empresas_exec)
        colunas = list(grade[0].keys()) if grade else ["Código"]
        ws.append(colunas)
        for linha in grade:
            ws.append([linha[c] for c in colunas])
        planilha = io.BytesIO()
        wb.save(planilha)
        zf.writestr(
            f"{nfgo_arquivos.PASTA_RAIZ}/Relatorio_RPA_NF_GO_{nfgo_arquivos.competencia_pasta(competencia)}.xlsx",
            planilha.getvalue(),
        )
    return buffer.getvalue()


def _botao_cancelar_execucao(execucao: dict, coluna, prefixo: str) -> None:
    """"⛔ Cancelar processamento": na fila, cancela na hora; rodando, o
    worker para na próxima etapa. O que não terminou vira erro "Cancelado
    pelo usuário" e o botão Reprocessar aparece (ver rpa_core)."""
    if execucao["status"] not in (rpa_core.STATUS_PENDENTE, rpa_core.STATUS_RODANDO):
        return
    if execucao.get("cancelar_solicitado"):
        coluna.caption("⛔ Cancelamento solicitado — o robô para na próxima etapa. Se o programa do PC "
                       "foi fechado ou travou, force o cancelamento:")
        if coluna.button("⛔ Forçar cancelamento", key=f"{prefixo}_forcar_{execucao['id']}"):
            qtd = rpa_core.forcar_cancelamento(execucao["id"], st.session_state.get("escritorio_id"))
            _flash("flash_rpa_hub", f"⛔ Execução cancelada ({qtd} consulta(s) marcadas) — use 🔁 Reprocessar quando quiser.")
            st.rerun()
        return
    if execucao["status"] == rpa_core.STATUS_RODANDO and execucao.get("modulo") in rpa_core.MODULOS_PC:
        contato = rpa_core.ultimo_contato_pc(st.session_state.get("escritorio_id"), execucao["modulo"])
        if contato is None or (contato["segundos"] or 0) > rpa_core.PC_SEM_CONTATO_S:
            coluna.warning("⚠️ O programa do PC parou de responder — se ele foi fechado, clique em "
                           "Cancelar processamento para liberar o Reprocessar.")
    if coluna.button("⛔ Cancelar processamento", key=f"{prefixo}_{execucao['id']}"):
        resultado = rpa_core.solicitar_cancelamento(execucao["id"], st.session_state.get("escritorio_id"))
        if resultado == "cancelada":
            _flash("flash_rpa_hub", "⛔ Execução cancelada — use 🔁 Reprocessar quando quiser rodar de novo.")
        elif resultado == "solicitado":
            _flash("flash_rpa_hub", "⛔ Cancelamento solicitado — o robô para na próxima etapa; depois use 🔁 Reprocessar.")
        st.rerun()


def _expander_execucao_nfgo(execucao: dict, emoji_status: dict) -> None:
    """Execução do RPA NF GO: grade Código | Empresa | CNPJ | IE |
    Competência | XML Entrada | XML Saída | Total | Status | Data/Hora e,
    por empresa, os ZIPs e o print (evidência) de cada consulta."""
    empresas_exec = rpa_core.listar_empresas(execucao["id"])
    concluidas = [e for e in empresas_exec if e["status"] == rpa_core.STATUS_CONCLUIDO]
    erros = sum(1 for e in empresas_exec if e["status"] == rpa_core.STATUS_ERRO)
    total_entrada = sum(e.get("qtd_xml") or 0 for e in empresas_exec if e.get("obrigacao") == "ENTRADA")
    total_saida = sum(e.get("qtd_xml") or 0 for e in empresas_exec if e.get("obrigacao") == "SAIDA")
    _competencia_exec = f" · competência {execucao['competencia']}" if execucao.get("competencia") else ""
    titulo = (
        f"{emoji_status.get(execucao['status'], '•')} Execução #{execucao['id']} — "
        f"{execucao['criado_em']:%d/%m/%Y %H:%M}{_competencia_exec} — "
        f"Entrada: {total_entrada} XMLs | Saída: {total_saida} XMLs | {execucao['status']}"
        f"{f' ({erros} consulta(s) com erro)' if erros else ''}"
    )
    with st.expander(titulo):
        _pasta = f" · Pasta: {execucao['pasta_destino']}" if execucao.get("pasta_destino") else ""
        st.caption(f"Planilha: {execucao['planilha_nome']} · Enviada por {execucao['criado_por']}{_pasta}")
        if concluidas and execucao.get("concluido_em"):
            _barra_retencao_arquivos(execucao["concluido_em"])
        _cols_acoes = st.columns(3)
        _botao_cancelar_execucao(execucao, _cols_acoes[2], "cancelar_nfgo")
        if any(e.get("xml_zip") or e.get("evidencia_png") for e in empresas_exec):
            _cols_acoes[0].download_button(
                "📦 Baixar tudo (.zip, uma pasta por empresa/competência/tipo)",
                _montar_zip_nfgo(execucao, empresas_exec),
                file_name=f"RPA NF GO {nfgo_arquivos.competencia_pasta(execucao.get('competencia') or '')}.zip",
                key=f"zip_nfgo_{execucao['id']}",
            )
        if erros:
            if _cols_acoes[1].button(
                f"🔁 Reprocessar {erros} consulta(s) com erro", key=f"reprocessar_nfgo_{execucao['id']}",
            ):
                qtd = rpa_core.reprocessar_falhas(execucao["id"])
                _flash(
                    "flash_rpa_hub",
                    f"✅ {qtd} consulta(s) voltaram para a fila — com o programa RPA NF GO aberto no PC "
                    "(\"Ficar aguardando o Hub\" marcado), começa sozinho em até 30 segundos.",
                )
                st.rerun()

        st.dataframe(_linhas_grade_nfgo(execucao, empresas_exec), use_container_width=True, hide_index=True)

        st.markdown("**Arquivos por empresa**")
        for empresa, entrada, saida in _agrupar_nfgo(empresas_exec):
            cols = st.columns([3, 2, 2])
            cols[0].write(nfgo_arquivos.nome_pasta_empresa(empresa["codigo"], empresa.get("razao_social") or ""))
            for col, tipo, linha in ((cols[1], "ENTRADA", entrada), (cols[2], "SAIDA", saida)):
                if linha is None:
                    continue
                col.caption(f"{tipo.title()}: {_status_linha_nfgo(linha)}")
                if linha.get("xml_zip"):
                    col.download_button(
                        f"⬇️ {linha['xml_zip_nome']}", bytes(linha["xml_zip"]), file_name=linha["xml_zip_nome"],
                        key=f"nfgo_zip_{linha['id']}",
                    )
                if linha.get("evidencia_png"):
                    with col.popover("🖼️ Print da consulta"):
                        st.image(bytes(linha["evidencia_png"]))
                if linha["status"] == rpa_core.STATUS_ERRO:
                    col.caption(f"⚠️ {linha.get('erro') or ''}")
                    if linha.get("screenshot_erro"):
                        with col.popover("🖼️ Tela do erro"):
                            st.image(bytes(linha["screenshot_erro"]))
                elif linha.get("observacao"):
                    col.caption(f"⚠️ {linha['observacao']}")


def _resumo_empresas_planilha(empresas: list[dict], eh_nfgo: bool) -> str:
    """'2 empresa(s)' - no RPA NF GO cada empresa vira duas linhas de fila
    (ENTRADA e SAIDA, ver rpa/sefazgo_nfe/planilha.py), então conta só as
    empresas e mostra as consultas à parte, em vez de len(empresas)."""
    if not eh_nfgo:
        return f"{len(empresas)} empresa(s)"
    qtd = len({(e["codigo"], e.get("inscricao_estadual")) for e in empresas})
    return f"{qtd} empresa(s) ({len(empresas)} consultas: Entrada e Saída)"


def _bloco_credenciais_certificado_senha(
    modulo_id: str, modulo_info: dict, escritorio_id: str, usuario: str, com_certificado: bool = True,
) -> dict | None:
    """Credenciais do RPA NF GO: CPF + senha do Acesso Restrito da SEFAZ-GO
    (obrigatório - é o que o formulário de login pede) e o certificado A1
    do escritório (opcional - apresentado no TLS quando o portal pedir, no
    lugar da janela "Selecionar certificado" do navegador). As duas ficam
    cifradas em rpa_credenciais (RPA_ENC_KEY), nunca em código. Retorna os
    metadados da credencial do portal (None = ainda não cadastrada)."""
    sistema_portal = modulo_info["sistema_credencial_portal"]
    sistema_cert = modulo_info["sistema_credencial"]

    st.subheader("Credenciais")
    cred_portal = rpa_core.tem_credencial(escritorio_id, sistema_portal)
    cred_cert = rpa_core.tem_credencial(escritorio_id, sistema_cert)
    col_portal, col_cert = st.columns(2) if com_certificado else (st.container(), None)
    with col_portal:
        if cred_portal:
            st.caption(
                f"✅ Acesso SEFAZ-GO — CPF {cred_portal['cnpj']} · "
                f"atualizado em {cred_portal['atualizado_em']:%d/%m/%Y %H:%M} por {cred_portal['atualizado_por']}"
            )
        else:
            st.caption("⚠️ CPF/senha do Acesso Restrito ainda não cadastrados (obrigatório).")
    if col_cert is not None:
        if cred_cert:
            col_cert.caption(
                f"✅ Certificado — titular {cred_cert['cnpj']} · "
                f"atualizado em {cred_cert['atualizado_em']:%d/%m/%Y %H:%M} por {cred_cert['atualizado_por']}"
            )
        else:
            col_cert.caption("ℹ️ Nenhum certificado cadastrado (opcional — só se o portal pedir).")

    with st.expander("Cadastrar / atualizar acesso SEFAZ-GO (CPF + senha)"):
        with st.form(f"credencial_portal_form_{modulo_id}", clear_on_submit=True):
            cpf_acesso = st.text_input("CPF de acesso (cadastro do escritório)")
            senha_acesso = st.text_input("Senha do Acesso Restrito", type="password")
            salvar_portal = st.form_submit_button("Salvar", type="primary")
        if salvar_portal:
            if not cpf_acesso.strip() or not senha_acesso:
                st.error("Informe CPF e senha.")
            else:
                cpf_limpo = re.sub(r"\D", "", cpf_acesso)
                rpa_core.salvar_credencial(escritorio_id, sistema_portal, cpf_limpo, senha_acesso, usuario)
                _flash("flash_rpa_hub", "✅ Acesso SEFAZ-GO salvo.")
                st.rerun()

    if not com_certificado:
        return cred_portal
    with st.expander("Cadastrar / atualizar certificado digital do escritório (A1)"):
        st.caption(
            "Só certificado A1 (arquivo .pfx/.p12) — A3 (token/cartão) não dá pra usar no "
            "worker do servidor. O robô apresenta este certificado quando o portal pedir, "
            "no lugar da janela de seleção de certificado do navegador."
        )
        with st.form(f"credencial_cert_form_{modulo_id}", clear_on_submit=True):
            titular_cert = st.text_input("CNPJ/CPF do titular do certificado")
            up_certificado = st.file_uploader("Certificado digital A1 (.pfx/.p12)", type=["pfx", "p12"])
            senha_cert = st.text_input("Senha do certificado", type="password")
            salvar_cert = st.form_submit_button("Salvar", type="primary")
        if salvar_cert:
            if not titular_cert.strip() or not up_certificado or not senha_cert:
                st.error("Informe o titular, o arquivo do certificado e a senha.")
            else:
                rpa_core.salvar_credencial_certificado(
                    escritorio_id, sistema_cert, titular_cert.strip(), up_certificado.getvalue(), senha_cert, usuario,
                )
                _flash("flash_rpa_hub", "✅ Certificado salvo.")
                st.rerun()
    return cred_portal


def _bloco_programa_nfgo(escritorio_id: str) -> None:
    """Download do programa do PC do RPA NF GO (attended_worker/gui_nfgo.py):
    a consulta de notas da SEFAZ-GO tem verificação da Cloudflare que não
    passa no navegador do servidor, então quem processa a fila é o Edge do
    escritório. O .exe é montado no deploy (job build-nfgo-exe)."""
    st.subheader("Programa do PC (processa a fila)")
    st.caption(
        "A consulta de notas da SEFAZ-GO só libera no navegador do escritório. Crie a execução "
        "aqui e processe pelo programa RPA NF GO: ele abre o Edge, confirma o certificado do "
        "escritório, faz a nova autenticação com o CPF/senha cadastrados abaixo e preenche a "
        "consulta. Os resultados (quantidades, print e ZIP) voltam para a grade abaixo."
    )
    _pasta_attended = Path(__file__).parent / "attended_worker"
    _exe = _pasta_attended / "dist" / "nfgo_attended.exe"
    if not _exe.exists():
        st.warning("O programa ainda não está disponível para download nesta versão do Hub.")
        return
    _zip = io.BytesIO()
    with zipfile.ZipFile(_zip, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.write(_exe, arcname="nfgo_attended.exe")
        _leia = _pasta_attended / "LEIA-ME_NFGO.txt"
        if _leia.exists():
            zf.write(_leia, arcname="LEIA-ME.txt")
    _contato = rpa_core.ultimo_contato_pc(escritorio_id, "sefazgo_nfe")
    if _contato and _contato["segundos"] is not None and _contato["segundos"] <= 90:
        st.success(f"🟢 Programa aberto no PC e aguardando a fila (último contato há {_contato['segundos']} s).")
    elif _contato:
        st.caption(f"⚪ Programa do PC sem contato desde {_contato['visto_em']:%d/%m/%Y %H:%M} — abra o RPA NF GO no PC para processar a fila.")
    _versao_arq = _pasta_attended / "VERSION_NFGO"
    _versao = _versao_arq.read_text(encoding="utf-8").strip() if _versao_arq.exists() else ""
    st.download_button(
        f"📥 Baixar programa RPA NF GO{f' v{_versao}' if _versao else ''} (.zip)",
        _zip.getvalue(), file_name="rpa_nf_go.zip", type="primary", key="download_programa_nfgo",
    )
    st.caption(
        "Extraia e dê duplo clique no .exe (cria o atalho \"RPA NF GO\" na Área de Trabalho e se "
        "atualiza sozinho). O Windows pode avisar 'aplicativo desconhecido' (SmartScreen): "
        "'Mais informações' → 'Executar assim mesmo'."
    )


def _tela_rpa_hub(modulos: list | None = None, titulo: str = "🤖 Hub de RPAs") -> None:
    """Hub de RPAs do escritório: cadastro de credenciais de procurador,
    upload de planilha e acompanhamento das execuções, por módulo (ISS Web
    é o primeiro — novos módulos só entram em rpa/registry.py, esta tela
    não muda). O processamento de verdade roda no worker separado
    (rpa_worker.py, container à parte com Playwright) — esta tela só
    enfileira em rpa_execucoes/rpa_empresas e mostra status/resultado.

    modulos: quais módulos esta tela mostra - padrão são os que não têm card
    próprio na home (sem 'app_home' no registry); o RPA NF GO passa só o dele.

    Escritório: qualquer usuário (inclusive comum) só vê/mexe nas
    credenciais, listas e execuções do PRÓPRIO escritório - só
    super_admin_global tem o seletor pra escolher QUALQUER escritório
    (mesmo padrão de _tela_historico/_tela_gerenciar_clientes), pra poder
    dar suporte sem precisar logar como o cliente."""
    st.title(titulo)
    if st.button("← Início"):
        st.session_state["tela"] = "home"
        st.rerun()
    st.divider()

    eh_global = st.session_state.get("papel_usuario") == auth.PAPEL_SUPER_GLOBAL
    if eh_global:
        escritorios = auth.carregar_escritorios()
        if not escritorios:
            st.info("Nenhum escritório cadastrado ainda.")
            return
        escritorio_id = st.selectbox(
            "Escritório",
            list(escritorios.keys()),
            format_func=lambda eid: escritorios.get(eid, {}).get("nome", eid),
            index=list(escritorios.keys()).index(st.session_state.get("escritorio_id"))
            if st.session_state.get("escritorio_id") in escritorios else 0,
        )
    else:
        escritorio_id = st.session_state.get("escritorio_id")
    usuario = st.session_state.get("usuario_logado")

    if modulos is None:
        modulos = [m for m, info in rpa_registry.MODULOS.items() if not info.get("app_home")]
    if len(modulos) == 1:
        modulo_id = modulos[0]
        st.caption(rpa_registry.MODULOS[modulo_id]["titulo"])
    else:
        modulo_id = st.selectbox(
            "Rotina", modulos,
            format_func=lambda m: rpa_registry.MODULOS[m]["titulo"],
        )
    modulo_info = rpa_registry.MODULOS[modulo_id]
    sistema = modulo_info["sistema_credencial"]

    if not modulo_info.get("automatizado", True):
        _tela_rpa_manual(modulo_id, modulo_info, escritorio_id, usuario)
        return

    _periodo_padrao = rpa_registry.preparar_periodo(modulo_id)
    _data_competencia = st.date_input(
        "Competência a executar", value=date(_periodo_padrao["ano"], _periodo_padrao["mes"], 1),
        key=f"competencia_{modulo_id}", format="DD/MM/YYYY",
        help="Escolha qualquer dia dentro do mês/ano desejado — só o mês e o ano importam, "
             "o dia é ignorado. Vem pré-preenchido com o mês anterior ao de hoje.",
    )
    competencia_escolhida = f"{_data_competencia.month:02d}/{_data_competencia.year}"

    tipo_auth = modulo_info.get("tipo_auth", "senha")
    if modulo_id == "sefazgo_nfe":
        # roda no programa do PC (Edge do escritório, certificado do
        # escritório) - não usa credencial guardada no Hub
        _bloco_programa_nfgo(escritorio_id)
        # CPF + senha do Acesso Restrito: o programa do PC usa na tela
        # "Este módulo requer nova autenticação" (o certificado fica no Windows)
        _bloco_credenciais_certificado_senha(modulo_id, modulo_info, escritorio_id, usuario, com_certificado=False)
        cred = True
    elif tipo_auth == "certificado_senha":
        cred = _bloco_credenciais_certificado_senha(modulo_id, modulo_info, escritorio_id, usuario)
    else:
        st.subheader("Credenciais do procurador")
        cred = rpa_core.tem_credencial(escritorio_id, sistema)
        if cred:
            st.caption(
                f"✅ Cadastrada — CNPJ {cred['cnpj']} · "
                f"atualizado em {cred['atualizado_em']:%d/%m/%Y %H:%M} por {cred['atualizado_por']}"
            )
        else:
            st.caption("⚠️ Nenhuma credencial cadastrada ainda para esta rotina.")

        with st.expander("Cadastrar / atualizar credencial"):
            if tipo_auth == "certificado":
                st.caption(
                    "Só certificado A1 (arquivo .pfx/.p12) — A3 é um token/leitor físico, "
                    "não dá pra usar no worker do servidor. Para A3, feche essa empresa pelo "
                    "navegador da própria máquina onde o leitor está instalado."
                )
                with st.form(f"credencial_form_{modulo_id}", clear_on_submit=True):
                    cnpj_proc = st.text_input("CPF/CNPJ do titular do certificado")
                    up_certificado = st.file_uploader("Certificado digital A1 (.pfx/.p12)", type=["pfx", "p12"])
                    senha_cert = st.text_input("Senha do certificado", type="password")
                    salvar_cred = st.form_submit_button("Salvar", type="primary")
                if salvar_cred:
                    if not cnpj_proc.strip() or not up_certificado or not senha_cert:
                        st.error("Informe CPF/CNPJ, o arquivo do certificado e a senha.")
                    else:
                        rpa_core.salvar_credencial_certificado(
                            escritorio_id, sistema, cnpj_proc.strip(), up_certificado.getvalue(), senha_cert, usuario,
                        )
                        _flash("flash_rpa_hub", "✅ Certificado salvo.")
                        st.rerun()
            else:
                with st.form(f"credencial_form_{modulo_id}", clear_on_submit=True):
                    cnpj_proc = st.text_input("CNPJ do procurador")
                    senha_proc = st.text_input("Senha do procurador", type="password")
                    salvar_cred = st.form_submit_button("Salvar", type="primary")
                if salvar_cred:
                    if not cnpj_proc.strip() or not senha_proc:
                        st.error("Informe CNPJ e senha.")
                    else:
                        rpa_core.salvar_credencial(escritorio_id, sistema, cnpj_proc.strip(), senha_proc, usuario)
                        _flash("flash_rpa_hub", "✅ Credencial salva.")
                        st.rerun()

    st.divider()
    _flash("flash_rpa_hub")

    st.subheader("Nova execução")
    _eh_nfgo = modulo_id == "sefazgo_nfe"
    if not cred:
        st.info("Cadastre a credencial do procurador acima antes de enviar uma planilha.")
    else:
        pasta_destino = ""
        if _eh_nfgo:
            # vem com a pasta da última execução (o escritório costuma usar sempre a mesma)
            _ultima_pasta = next(
                (e["pasta_destino"] for e in rpa_core.listar_execucoes(escritorio_id, modulo_id) if e.get("pasta_destino")),
                r"C:\RPA NF GO",
            )
            pasta_destino = st.text_input(
                "Pasta de destino dos XMLs",
                value=_ultima_pasta,
                key=f"pasta_destino_{modulo_id}",
                help="Referência da pasta usada no PC do escritório (no programa RPA NF GO você "
                     "escolhe a pasta de verdade). Dentro dela: RPA NF GO / CÓDIGO - EMPRESA / MMAAAA / "
                     "ENTRADA e SAIDA, com ENTRADA_MMAAAA.zip, SAIDA_MMAAAA.zip e o print de cada consulta. "
                     "Os mesmos arquivos também ficam no \"Baixar tudo (.zip)\" abaixo.",
            )
        st.caption("Colunas esperadas: " + " · ".join(modulo_info["colunas_planilha"]))
        st.download_button(
            "📥 Baixar planilha modelo (.xlsx)", _planilha_modelo(modulo_info),
            file_name=f"modelo_empresas_{modulo_id}.xlsx", key=f"modelo_{modulo_id}",
        )
        up_planilha = st.file_uploader(
            "Planilha de empresas (.xlsx)", type=["xlsx"], key=f"up_planilha_{modulo_id}",
        )
        if up_planilha:
            conteudo = up_planilha.getvalue()
            try:
                empresas = rpa_registry.ler_empresas(modulo_id, conteudo)
            except Exception as exc:
                st.error(str(exc))
            else:
                st.success(_resumo_empresas_planilha(empresas, _eh_nfgo) + " encontrada(s) na planilha para esta rotina.")
                with st.expander("Ver empresas identificadas"):
                    if _eh_nfgo:
                        st.dataframe(
                            [
                                {"Código": e["codigo"], "Empresa": e["razao_social"], "CNPJ": e["cnpj_cpf"],
                                 "Inscrição Estadual": e["inscricao_estadual"]}
                                for e in empresas if e["obrigacao"] == "ENTRADA"
                            ],
                            use_container_width=True, hide_index=True,
                        )
                    else:
                        st.dataframe(
                            [
                                {"Código": e["codigo"], "CNPJ/CPF": e["cnpj_cpf"], "Obrigação": e["obrigacao"]}
                                for e in empresas
                            ],
                            use_container_width=True, hide_index=True,
                        )
                if st.button(
                    "🚀 Criar execução (processar no programa do PC)" if _eh_nfgo else "🚀 Iniciar processamento",
                    type="primary",
                ):
                    execucao_id = rpa_core.criar_execucao(
                        escritorio_id, modulo_id, conteudo, up_planilha.name, usuario, empresas,
                        competencia_escolhida, pasta_destino=pasta_destino.strip(),
                    )
                    _flash(
                        "flash_rpa_hub",
                        f"✅ Execução #{execucao_id} criada — competência {competencia_escolhida}, "
                        f"{_resumo_empresas_planilha(empresas, _eh_nfgo)} na fila.",
                    )
                    st.rerun()

    st.divider()
    st.subheader("Execuções")
    # lista de execuções num fragmento: o "Atualizar status" (e a
    # atualização automática enquanto há execução ativa) recarrega só esta
    # parte, sem refazer a tela inteira
    _ativa_agora = any(
        e["status"] in (rpa_core.STATUS_PENDENTE, rpa_core.STATUS_RODANDO)
        for e in rpa_core.listar_execucoes(escritorio_id, modulo_id)
    )

    @st.fragment(run_every=15 if _ativa_agora else None)
    def _lista_execucoes() -> None:
        rpa_core.limpar_arquivos_vencidos()
        execucoes = rpa_core.listar_execucoes(escritorio_id, modulo_id)
        _ativa = any(e["status"] in (rpa_core.STATUS_PENDENTE, rpa_core.STATUS_RODANDO) for e in execucoes)
        if not execucoes:
            st.info("Nenhuma execução ainda para esta rotina.")
            return

        # o clique já recarrega só este fragmento (lista de execuções)
        st.button("🔄 Atualizar status", key=f"atualizar_status_{modulo_id}")
        if _ativa:
            st.caption("🔄 Atualizando sozinho a cada 15 segundos enquanto há execução na fila ou rodando.")

        emoji_status = {"PENDENTE": "⏳", "RODANDO": "🔄", "CONCLUIDO": "✅", "ERRO": "❌"}
        if _eh_nfgo:
            for execucao in execucoes:
                _expander_execucao_nfgo(execucao, emoji_status)
            return
        for execucao in execucoes:
            empresas_exec = rpa_core.listar_empresas(execucao["id"])
            concluidas = sum(1 for e in empresas_exec if e["status"] == rpa_core.STATUS_CONCLUIDO)
            erros = sum(1 for e in empresas_exec if e["status"] == rpa_core.STATUS_ERRO)
            _competencia_exec = f" · competência {execucao['competencia']}" if execucao.get("competencia") else ""
            titulo = (
                f"{emoji_status.get(execucao['status'], '•')} Execução #{execucao['id']} — "
                f"{execucao['criado_em']:%d/%m/%Y %H:%M}{_competencia_exec} — {execucao['status']} "
                f"({concluidas}/{len(empresas_exec)} concluídas, {erros} erro(s))"
            )
            with st.expander(titulo):
                st.caption(f"Planilha: {execucao['planilha_nome']} · Enviada por {execucao['criado_por']}")
                if concluidas and execucao.get("concluido_em"):
                    _barra_retencao_arquivos(execucao["concluido_em"])
                _cols_acoes = st.columns(3)
                _botao_cancelar_execucao(execucao, _cols_acoes[2], "cancelar")
                if concluidas:
                    _pasta_competencia_zip = (execucao.get("competencia") or "sem-competencia").replace("/", "")
                    _cols_acoes[0].download_button(
                        "📦 Baixar tudo (.zip, uma pasta por empresa/competência)",
                        _montar_zip_execucao(execucao, empresas_exec),
                        file_name=f"DMS-XML {_pasta_competencia_zip}.zip",
                        key=f"zip_execucao_{execucao['id']}",
                    )
                if erros:
                    if _cols_acoes[1].button(
                        f"🔁 Reprocessar {erros} empresa(s) com erro", key=f"reprocessar_{execucao['id']}",
                    ):
                        qtd = rpa_core.reprocessar_falhas(execucao["id"])
                        _flash("flash_rpa_hub", f"✅ {qtd} empresa(s) voltaram para a fila — o worker processa em instantes.")
                        st.rerun()
                for empresa in empresas_exec:
                    cols = st.columns([1, 2, 1, 1, 2])
                    cols[0].write(empresa["codigo"])
                    cols[1].write(empresa["cnpj_cpf"])
                    cols[2].write(empresa["obrigacao"])
                    cols[3].write(empresa["status"])
                    if empresa["status"] == rpa_core.STATUS_CONCLUIDO and empresa["pdf"]:
                        cols[4].download_button(
                            "⬇️ PDF", bytes(empresa["pdf"]), file_name=empresa["pdf_nome"],
                            key=f"pdf_{empresa['id']}",
                        )
                        if empresa.get("xml_zip"):
                            _qtd_xml = _contar_arquivos_zip(bytes(empresa["xml_zip"]))
                            cols[4].download_button(
                                f"⬇️ XML ({_qtd_xml})", bytes(empresa["xml_zip"]), file_name=empresa["xml_zip_nome"],
                                key=f"xml_{empresa['id']}",
                            )
                    elif empresa["status"] == rpa_core.STATUS_ERRO:
                        cols[4].caption(f"⚠️ {empresa['erro']}")
                        if empresa.get("screenshot_erro"):
                            with cols[4].popover("🖼️ Ver tela do erro"):
                                st.image(bytes(empresa["screenshot_erro"]))
                    else:
                        cols[4].write("—")

    _lista_execucoes()


_APPS_HOME = [
    {
        "id": "conciliacao",
        "icone": "🏦",
        "titulo": "Conciliação Bancária Automática",
        "descricao": "Extrato/fluxo de caixa × razão contábil × balancete → espelho e arquivo de importação Domínio.",
        "tela": "conciliacao",
    },
    {
        "id": "rpa_hub",
        "icone": "🤖",
        "titulo": "RPA — Fechamento REST/DMS",
        "descricao": "Fechamento mensal de REST e DMS no ISS Web.",
        "tela": "rpa_hub",
    },
    {
        "id": "rpa_folha",
        "icone": "📋",
        "titulo": "RPA — Folha de Pagamento",
        "descricao": "Fechamento da folha no Domínio Folha. Ainda roda por automação assistida, sem tela própria aqui.",
        "tela": None,
    },
    {
        "id": "rpa_nfgo",
        "icone": "🧾",
        "titulo": "RPA NF GO",
        "descricao": "Download mensal dos XMLs de NF-e (Entrada e Saída) na SEFAZ-GO, com quantidade de notas e print da consulta.",
        "tela": "rpa_nfgo",
    },
]


_CARDS_POR_LINHA = 3


def _apps_permitidos_efetivos(escritorio_id: str, usuario: str) -> set:
    """IDs de app (_APPS_HOME) visíveis para este usuário: interseção entre
    o teto do escritório (todos, se ele não restringiu nada) e o refino do
    próprio usuário (todos os do escritório, se ele não restringiu nada) -
    configurado em Gerenciar Escritórios/Usuários. super_admin_global nunca
    é restringido (é quem administra o hub inteiro)."""
    if st.session_state.get("papel_usuario") == auth.PAPEL_SUPER_GLOBAL:
        return {a["id"] for a in _APPS_HOME}

    todos = {a["id"] for a in _APPS_HOME}
    escritorio = auth.carregar_escritorios().get(escritorio_id, {})
    apps_escritorio = set(escritorio.get("apps_permitidos") or todos) & todos

    usuarios = auth.carregar_usuarios()
    apps_usuario = set(usuarios.get(usuario, {}).get("apps_permitidos") or apps_escritorio)

    return apps_escritorio & apps_usuario


def _tela_home() -> None:
    """Tela inicial: um icone por aplicativo do escritorio. Cada app novo
    (proxima automacao) so precisa de uma entrada em _APPS_HOME - nao mexe
    no roteamento das telas que ja existem. A sidebar daqui e a unica
    dona de Gerenciar Escritorios/Usuarios - a tela de Conciliacao
    Bancaria nao mostra mais esses dois (sao administracao do hub, nao
    algo especifico daquele app)."""
    papel_usuario = st.session_state.get("papel_usuario")
    with st.sidebar:
        st.caption(f"👤 {st.session_state.get('nome_usuario')} · {papel_usuario}")
        st.caption(f"🏢 {st.session_state.get('escritorio_nome')}")
        csb1, csb2 = st.columns(2)
        with csb1:
            if st.button("Trocar senha", use_container_width=True, key="home_trocar_senha"):
                st.session_state["mostrar_trocar_senha"] = True
                st.rerun()
        with csb2:
            if st.button("Sair", use_container_width=True, key="home_sair"):
                _fazer_logout()
        if papel_usuario in (auth.PAPEL_SUPER_GLOBAL, auth.PAPEL_ADMIN_ESCRITORIO):
            st.divider()
            st.caption("Administração")
            if papel_usuario == auth.PAPEL_SUPER_GLOBAL:
                if st.button("🌐 Gerenciar Escritórios", use_container_width=True, key="home_gerenciar_escritorios"):
                    st.session_state["tela"] = "gerenciar_escritorios"
                    st.rerun()
            if st.button("👥 Gerenciar Usuários", use_container_width=True, key="home_gerenciar_usuarios"):
                st.session_state["tela"] = "gerenciar_usuarios"
                st.rerun()

    st.markdown(
        "<div style='text-align:center;font-size:2.25rem;font-weight:700;"
        "margin-bottom:0.25rem;'>👋 Bem-vindo(a)</div>"
        "<div style='text-align:center;color:rgba(250,250,250,0.6);"
        "margin-bottom:2rem;'>Escolha um aplicativo para começar.</div>",
        unsafe_allow_html=True,
    )

    _ids_visiveis = _apps_permitidos_efetivos(
        st.session_state.get("escritorio_id"), st.session_state.get("usuario_logado"),
    )
    _apps_visiveis = [a for a in _APPS_HOME if a["id"] in _ids_visiveis]
    if not _apps_visiveis:
        st.info("Nenhum aplicativo liberado para o seu usuário ainda — fale com o administrador do seu escritório.")
        return

    # grade de _CARDS_POR_LINHA: app novo entra na linha de baixo com o
    # mesmo tamanho de card, em vez de espremer todos numa linha só (a
    # última linha incompleta fica alinhada à esquerda, colunas vazias).
    for inicio in range(0, len(_apps_visiveis), _CARDS_POR_LINHA):
        colunas = st.columns(_CARDS_POR_LINHA)
        for coluna, app in zip(colunas, _apps_visiveis[inicio:inicio + _CARDS_POR_LINHA]):
            _card_app(coluna, app)


def _card_app(coluna, app: dict) -> None:
    """Um card da home (ícone, título, descrição e botão Abrir/Em breve)."""
    with coluna:
        with st.container(border=True):
            # titulo+descricao num min-height fixo (em vez de
            # st.caption separado): garante que os 3 cards tenham a
            # mesma altura e o botao "Abrir" comece sempre na mesma
            # posicao, mesmo com textos de tamanhos diferentes.
            st.markdown(
                f"<div style='text-align:center;font-size:3rem;margin-bottom:0.4rem;'>{app['icone']}</div>"
                f"<div style='min-height:150px;'>"
                f"<div style='text-align:center;font-weight:600;margin-bottom:0.3rem;'>{app['titulo']}</div>"
                f"<div style='text-align:center;color:rgba(250,250,250,0.6);font-size:0.875rem;'>"
                f"{app['descricao']}</div>"
                f"</div>",
                unsafe_allow_html=True,
            )
            if app["tela"]:
                if st.button("Abrir", key=f"home_abrir_{app['tela']}",
                             use_container_width=True, type="primary"):
                    st.session_state["tela"] = app["tela"]
                    st.rerun()
            else:
                st.button("Em breve", key="home_abrir_rpa_folha",
                          use_container_width=True, disabled=True)


_renderizar_cookie_pendente()

if "usuario_logado" not in st.session_state:
    # antes de exigir login de novo: a pagina pode ter sido so atualizada
    # (F5), que abre uma conexao Streamlit nova e zera o session_state -
    # se o navegador ainda tiver um cookie de sessao valido, reloga sem
    # pedir usuario/senha de novo.
    _token_cookie = st.context.cookies.get(_COOKIE_SESSAO, "")
    _dados_cookie = auth.validar_sessao(_token_cookie) if _token_cookie else None
    if _dados_cookie:
        _popular_sessao(_dados_cookie, token=_token_cookie)
        st.rerun()

if "usuario_logado" not in st.session_state:
    _tela_login()
    st.stop()

if st.session_state.get("deve_trocar_senha"):
    _tela_trocar_senha(obrigatoria=True)
    st.stop()

if st.session_state.get("mostrar_trocar_senha"):
    _tela_trocar_senha(obrigatoria=False)
    st.stop()

# tela atual também na URL (?tela=...): F5 ou uma sessão nova do Streamlit
# (reconexão, Hub reiniciado no deploy) voltava sempre pra tela inicial
_TELAS_URL = {"home", "rpa_hub", "rpa_nfgo", "conciliacao", "historico", "gerenciar_clientes",
              "gerenciar_escritorios", "gerenciar_usuarios"}
if "tela" not in st.session_state:
    _tela_url = st.query_params.get("tela", "home")
    st.session_state["tela"] = _tela_url if _tela_url in _TELAS_URL else "home"
if st.session_state["tela"] in _TELAS_URL and st.query_params.get("tela") != st.session_state["tela"]:
    st.query_params["tela"] = st.session_state["tela"]

if st.session_state.get("tela") == "home":
    _tela_home()
    st.stop()

if st.session_state.get("tela") == "gerenciar_escritorios":
    if st.session_state.get("papel_usuario") != auth.PAPEL_SUPER_GLOBAL:
        st.session_state["tela"] = "home"
    else:
        _tela_gerenciar_escritorios()
        st.stop()

if st.session_state.get("tela") == "gerenciar_usuarios":
    if st.session_state.get("papel_usuario") not in (auth.PAPEL_SUPER_GLOBAL, auth.PAPEL_ADMIN_ESCRITORIO):
        st.session_state["tela"] = "home"
    else:
        _tela_gerenciar_usuarios()
        st.stop()

if st.session_state.get("tela") == "historico":
    # qualquer usuario logado pode ver (inclusive usuario comum) - sempre
    # restrito ao proprio escritorio, ja garantido dentro de _tela_historico.
    _tela_historico()
    st.stop()

if st.session_state.get("tela") == "gerenciar_clientes":
    _tela_gerenciar_clientes()
    st.stop()

if st.session_state.get("tela") == "rpa_hub":
    if "rpa_hub" not in _apps_permitidos_efetivos(
        st.session_state.get("escritorio_id"), st.session_state.get("usuario_logado")
    ):
        st.session_state["tela"] = "home"
        st.rerun()
    _tela_rpa_hub()
    st.stop()

if st.session_state.get("tela") == "rpa_nfgo":
    if "rpa_nfgo" not in _apps_permitidos_efetivos(
        st.session_state.get("escritorio_id"), st.session_state.get("usuario_logado")
    ):
        st.session_state["tela"] = "home"
        st.rerun()
    _tela_rpa_hub(
        modulos=[m for m, info in rpa_registry.MODULOS.items() if info.get("app_home") == "rpa_nfgo"],
        titulo="🧾 RPA NF GO",
    )
    st.stop()

if st.session_state.get("tela") == "conciliacao":
    if "conciliacao" not in _apps_permitidos_efetivos(
        st.session_state.get("escritorio_id"), st.session_state.get("usuario_logado")
    ):
        st.session_state["tela"] = "home"
        st.rerun()
    # sem st.stop() no caso permitido: cai no corpo principal do arquivo
    # (a tela de Conciliacao Bancaria sempre foi o "resto do script", sem
    # funcao propria - so ganhou esse guard explicito aqui pra checar
    # permissao antes de renderizar).

with st.sidebar:
    st.caption(f"👤 {st.session_state.get('nome_usuario')} · {st.session_state.get('papel_usuario')}")
    st.caption(f"🏢 {st.session_state.get('escritorio_nome')}")
    csb1, csb2 = st.columns(2)
    with csb1:
        if st.button("Trocar senha", use_container_width=True):
            st.session_state["mostrar_trocar_senha"] = True
            st.rerun()
    with csb2:
        if st.button("Sair", use_container_width=True):
            _fazer_logout()
    if st.button("🏠 Início", use_container_width=True):
        st.session_state["tela"] = "home"
        st.rerun()
    if st.button("🏢 Gerenciar Clientes", use_container_width=True):
        st.session_state["tela"] = "gerenciar_clientes"
        st.rerun()
    if st.button("📜 Histórico de Lançamentos", use_container_width=True):
        st.session_state["tela"] = "historico"
        st.rerun()
    st.divider()

def _resetar_empresa() -> None:
    """Limpa os arquivos e parametros preenchidos, pra comecar a
    conciliacao de outra empresa do zero - chamada tanto pelo botao do
    topo quanto pelo de baixo, perto dos downloads."""
    _chaves_reset_empresa = [
        "forcar_escolha_conta",
        "sel_conta_saida_padrao", "sel_conta_entrada_padrao", "filtro_pendencias",
        "sel_conta_grupo_filtrado", "historico_ultima_assinatura", "pend_contas_manuais",
        "pagina_pendencias_individual",
    ]
    for _chave in _chaves_reset_empresa:
        st.session_state.pop(_chave, None)
    # selectboxes individuais de tratativa de pendencia (chave dinamica
    # por lancamento) - limpa todas de uma vez, nao da pra listar antes.
    for _chave in list(st.session_state.keys()):
        if _chave.startswith("sel_ci_"):
            del st.session_state[_chave]
    # os 3 file_uploader tem uma peculiaridade do Streamlit: apagar a
    # chave do session_state sozinho nao limpa visualmente o arquivo ja
    # selecionado - precisa trocar a propria key do widget (por isso o
    # contador "reset_seq" usado no key= deles mais abaixo).
    st.session_state["reset_seq"] = st.session_state.get("reset_seq", 0) + 1
    # limpa o cache de processamento (@st.cache_data) tambem - garante que
    # a proxima conciliacao roda 100% do zero, sem reaproveitar nenhum
    # resultado (balancete/conta banco/conciliacao) calculado antes.
    st.cache_data.clear()


def _resetar_lancamento() -> None:
    """Limpa os arquivos e pendencias do lancamento atual (extrato,
    balancete, competencia, saldos) pra rodar a proxima competencia -
    mas MANTEM a empresa (cliente) e o modo de arquivo selecionados,
    ja que normalmente e a mesma empresa que vai lancar o mes seguinte
    (diferente do botao "Novo lançamento" do topo, que zera tudo,
    inclusive a empresa, pra comecar outra do zero)."""
    _seq_atual = st.session_state.get("reset_seq", 0)
    _cliente_atual = st.session_state.get(f"sel_cliente_empresa_{_seq_atual}")
    _modo_atual = st.session_state.get(f"radio_modo_arquivos_{_seq_atual}")
    _resetar_empresa()
    _seq_novo = st.session_state.get("reset_seq", 0)
    if _cliente_atual:
        st.session_state[f"sel_cliente_empresa_{_seq_novo}"] = _cliente_atual
    if _modo_atual:
        st.session_state[f"radio_modo_arquivos_{_seq_novo}"] = _modo_atual


_titulo_col, _reset_col = st.columns([5, 1.4])
with _titulo_col:
    st.title("🏦 Conciliação Bancária Automatizada")
    st.caption(f"{st.session_state.get('escritorio_nome')} · "
               "Extrato/fluxo de caixa × razão contábil × balancete → espelho + arquivo de importação Domínio")
with _reset_col:
    st.write("")
    if st.button(
        "🆕 Novo lançamento", use_container_width=True, key="btn_nova_empresa_topo",
        help="Limpa arquivos, parâmetros e cache, pra começar a conciliação de outra empresa do zero.",
    ):
        _resetar_empresa()
        st.rerun()

# Alem dos file_uploader, o selectbox de empresa e o radio de modo de
# arquivo tambem sofrem da mesma peculiaridade do Streamlit (apagar a
# chave do session_state sozinho nao reseta visualmente o valor ja
# escolhido) - por isso usam o mesmo sufixo "reset_seq" no key= deles.
_reset_seq = st.session_state.get("reset_seq", 0)

# ---------------------------------------------------------------------------
# Deteccao automatica de arquivos numa pasta
# ---------------------------------------------------------------------------

_PADROES = {
    "extrato": ["extrato", "fluxo", "movimento", "ofx"],
    "razao": ["razao", "razão", "livro razao"],
    "balancete": ["balancete", "balanco", "balanço"],
}


def _identificar_arquivos_na_pasta(pasta: Path) -> dict:
    """Varre a pasta e tenta identificar extrato/razao/balancete pelo nome.
    Retorna dict {categoria: [candidatos]} - nao decide sozinho quando ha
    mais de um candidato, deixa o usuario escolher na interface."""
    candidatos = {k: [] for k in _PADROES}
    exts_validas = {".ofx", ".csv", ".pdf"}
    for f in pasta.iterdir():
        if not f.is_file() or f.suffix.lower() not in exts_validas:
            continue
        nome = f.name.lower()
        for categoria, termos in _PADROES.items():
            if any(t in nome for t in termos):
                candidatos[categoria].append(f)
    return candidatos


# ---------------------------------------------------------------------------
# Sidebar - parametros gerais
# ---------------------------------------------------------------------------

with st.sidebar:
    st.header("Parâmetros")

    _escritorio_id_atual = st.session_state.get("escritorio_id")
    _lista_clientes = clientes.listar_clientes(_escritorio_id_atual)
    _flash("flash_cliente_salvo")

    empresa_codigo: Optional[str] = None
    empresa_nome_selecionado: Optional[str] = None
    if _lista_clientes:
        _opcoes_cliente = ["(selecione)"] + [f"{c['nome']} — {c['codigo_dominio']}" for c in _lista_clientes]
        _escolha_cliente = st.selectbox(
            "Empresa (cliente) *", _opcoes_cliente, key=f"sel_cliente_empresa_{_reset_seq}",
        )
        if _escolha_cliente != "(selecione)":
            _cliente_sel = _lista_clientes[_opcoes_cliente.index(_escolha_cliente) - 1]
            empresa_codigo = _cliente_sel["codigo_dominio"]
            empresa_nome_selecionado = _cliente_sel["nome"]
    else:
        st.info("Nenhum cliente cadastrado ainda — cadastre um abaixo.")

    with st.expander("➕ Cadastrar novo cliente"):
        with st.form("novo_cliente_form", clear_on_submit=True):
            _cliente_nome_novo = st.text_input("Nome do cliente")
            _cliente_codigo_novo = st.text_input("Código da empresa no Domínio")
            _salvar_cliente = st.form_submit_button("Cadastrar cliente", type="primary")
        if _salvar_cliente:
            if not _cliente_nome_novo.strip() or not _cliente_codigo_novo.strip():
                st.error("Informe nome e código da empresa.")
            else:
                _ok, _msg = clientes.criar_cliente(
                    _escritorio_id_atual, _cliente_nome_novo.strip(), _cliente_codigo_novo.strip()
                )
                if _ok:
                    _flash("flash_cliente_salvo", f"✅ {_msg}")
                    st.rerun()
                else:
                    st.error(_msg)

        if _lista_clientes and st.session_state.get("papel_usuario") in (
            auth.PAPEL_SUPER_GLOBAL, auth.PAPEL_ADMIN_ESCRITORIO,
        ):
            st.caption("Remover cliente cadastrado:")
            for _c in _lista_clientes:
                _cc1, _cc2 = st.columns([4, 1])
                with _cc1:
                    st.caption(f"{_c['nome']} — {_c['codigo_dominio']}")
                with _cc2:
                    if st.button("🗑️", key=f"remover_cliente_{_c['id']}"):
                        clientes.remover_cliente(_c["id"], _escritorio_id_atual)
                        _flash("flash_cliente_salvo", f"✅ Cliente '{_c['nome']}' removido.")
                        st.rerun()

    competencia = st.text_input(
        "Competência (MM-AAAA)",
        value="",
        placeholder="Ex: 01-2026 — deixe em branco p/ detectar do balancete",
        help="Se deixado em branco, a competência é detectada automaticamente a partir do texto do balancete informado.",
        key=f"input_competencia_{_reset_seq}",
    )

    col1, col2 = st.columns(2)
    with col1:
        saldo_inicial_razao = st.number_input(
            "Saldo inicial razão (R$)", value=0.0, step=0.01, format="%.2f",
            key=f"input_saldo_razao_{_reset_seq}",
        )
    with col2:
        saldo_inicial_extrato = st.number_input(
            "Saldo inicial extrato (R$)", value=0.0, step=0.01, format="%.2f",
            key=f"input_saldo_extrato_{_reset_seq}",
        )

    with st.expander("Avançado"):
        dias_tolerancia = st.number_input(
            "Dias de tolerância no match (cheques)", value=0, min_value=0, max_value=15,
            key=f"input_dias_tolerancia_{_reset_seq}",
        )
        cod_historico = st.text_input(
            "Código de histórico padrão Domínio", value="",
            help="Confirme contra a tabela de históricos do escritório antes de importar.",
            key=f"input_cod_historico_{_reset_seq}",
        )

st.divider()

# ---------------------------------------------------------------------------
# Escolha de origem dos arquivos: upload ou pasta
# ---------------------------------------------------------------------------

modo = st.radio(
    "Como fornecer os arquivos?", ["Upload de arquivos", "Caminho de uma pasta"],
    horizontal=True, key=f"radio_modo_arquivos_{_reset_seq}",
)

extrato_path: Optional[str] = None
razao_path: Optional[str] = None
balancete_path: Optional[str] = None

if "tmpdir" not in st.session_state:
    st.session_state.tmpdir = Path(tempfile.mkdtemp(prefix="conc_upload_"))
_tmpdir = st.session_state.tmpdir
_reset_seq = st.session_state.get("reset_seq", 0)


def _salvar_upload(uploaded_file, destino_nome: str) -> str:
    destino = _tmpdir / destino_nome
    destino.write_bytes(uploaded_file.getvalue())
    return str(destino)


if modo == "Upload de arquivos":
    c1, c2, c3 = st.columns(3)
    with c1:
        up_extrato = st.file_uploader(
            "Extrato / fluxo de caixa", type=["ofx", "csv", "pdf"], key=f"up_extrato_{_reset_seq}",
        )
    with c2:
        up_razao = st.file_uploader(
            "Razão contábil da conta (opcional)", type=["csv", "pdf"], key=f"up_razao_{_reset_seq}",
        )
    with c3:
        up_balancete = st.file_uploader(
            "Balancete", type=["csv", "pdf"], key=f"up_balancete_{_reset_seq}",
        )

    if up_extrato:
        extrato_path = _salvar_upload(up_extrato, f"extrato{Path(up_extrato.name).suffix}")
    if up_razao:
        razao_path = _salvar_upload(up_razao, f"razao{Path(up_razao.name).suffix}")
    if up_balancete:
        balancete_path = _salvar_upload(up_balancete, f"balancete{Path(up_balancete.name).suffix}")

else:
    st.warning(
        "⚠️ Esse caminho é lido no **servidor** onde o app está rodando (não no seu computador). "
        "Como este sistema fica hospedado na nuvem (hub.redeg7.com), ele não enxerga pastas do seu "
        "PC — use \"Upload de arquivos\" acima para enviar os arquivos diretamente daqui. Esta opção "
        "só funciona quando o app roda localmente na própria máquina onde os arquivos estão."
    )
    pasta_str = st.text_input(
        "Caminho da pasta com os arquivos", placeholder=r"C:\Clientes\EmpresaX\2026-01",
        key=f"input_pasta_caminho_{_reset_seq}",
    )
    if pasta_str:
        pasta = Path(pasta_str)
        if not pasta.is_dir():
            st.error(f"Pasta não encontrada no servidor: {pasta_str} — lembre-se que este caminho "
                     "precisa existir na máquina onde o app está rodando, não no seu computador.")
        else:
            candidatos = _identificar_arquivos_na_pasta(pasta)
            st.write("Arquivos identificados automaticamente (confira/ajuste antes de rodar):")
            c1, c2, c3 = st.columns(3)

            def _seletor(col, categoria, label):
                opcoes = candidatos[categoria] or list(pasta.glob("*"))
                opcoes_validas = [f for f in opcoes if f.is_file() and f.suffix.lower() in (".ofx", ".csv", ".pdf")]
                nomes = ["(selecione)"] + [f.name for f in opcoes_validas]
                with col:
                    escolha = st.selectbox(label, nomes, key=f"sel_{categoria}_{_reset_seq}")
                if escolha != "(selecione)":
                    return str(pasta / escolha)
                return None

            extrato_path = _seletor(c1, "extrato", "Extrato")
            razao_path = _seletor(c2, "razao", "Razão (opcional)")
            balancete_path = _seletor(c3, "balancete", "Balancete")

# ---------------------------------------------------------------------------
# Conta banco: identificada automaticamente cruzando extrato x balancete.
# O usuario nao "lanca" nada aqui - so confirma se o app pedir, quando a
# identificacao automatica fica ambigua.
# ---------------------------------------------------------------------------


@st.cache_data(show_spinner=False)
def _parse_balancete_cached(path: str, _mtime: float):
    return cb.parse_balancete(path)


@st.cache_data(show_spinner=False)
def _detectar_conta_banco_cached(extrato_path: str, _extrato_mtime: float,
                                  balancete_path: str, _balancete_mtime: float):
    contas = _parse_balancete_cached(balancete_path, _balancete_mtime)
    return cb.detectar_conta_banco(extrato_path, contas)


conta_contabil: Optional[str] = None

if balancete_path:
    try:
        contas_balancete = _parse_balancete_cached(balancete_path, Path(balancete_path).stat().st_mtime)
    except Exception as exc:
        st.error(f"Erro ao ler o balancete: {exc}")
        contas_balancete = []

    _contas_nome_ilegivel = [
        c for c in contas_balancete if c.nome.startswith(cb.NOME_ILEGIVEL_BALANCETE_PREFIXO)
    ]
    if _contas_nome_ilegivel:
        with st.expander(
            f"⚠️ {len(_contas_nome_ilegivel)} conta(s) do balancete com nome ilegível "
            "(código recuperado, mas o nome não — clique para ver os códigos)"
        ):
            st.caption(
                "Nome comprido demais ficou sobreposto ao código no PDF de origem — o código foi "
                "recuperado (aparece como \"não identificado\" na lista de contas), mas confira o "
                "nome real no balancete original pela classificação antes de vincular:"
            )
            st.write(", ".join(f"{c.codigo} ({c.classificacao})" for c in _contas_nome_ilegivel))

    if contas_balancete:
        conta_auto = None
        if extrato_path:
            conta_auto = _detectar_conta_banco_cached(
                extrato_path, Path(extrato_path).stat().st_mtime,
                balancete_path, Path(balancete_path).stat().st_mtime,
            )

        if conta_auto and not st.session_state.get("forcar_escolha_conta"):
            conta_contabil = conta_auto.nome
            st.success(f"Conta banco identificada automaticamente a partir do extrato: "
                       f"{conta_auto.codigo} — {conta_auto.nome}")
            if st.button("Não é essa conta — escolher manualmente"):
                st.session_state["forcar_escolha_conta"] = True
                st.rerun()
        else:
            if not extrato_path:
                st.info("Informe o extrato para o app tentar identificar a conta banco automaticamente "
                         "(cruzando o banco do extrato com o balancete).")
            else:
                st.warning("Não consegui identificar a conta banco automaticamente a partir do extrato — "
                            "confirme qual é, abaixo (o balancete continua sendo a fonte de verdade).")
            provaveis = cb.contas_provaveis_banco(contas_balancete)
            mostrar_todas = st.checkbox(
                "Mostrar todas as contas do balancete (em vez de só as prováveis de banco/caixa)",
                value=not provaveis,
            )
            opcoes = contas_balancete if (mostrar_todas or not provaveis) else provaveis
            rotulos = ["(selecione)"] + [f"{c.codigo} — {c.nome}" for c in opcoes]
            escolha = st.selectbox(
                "Conta banco a conciliar (lida do balancete) *", rotulos, key=f"sel_conta_banco_{_reset_seq}",
            )
            if escolha != "(selecione)":
                conta_contabil = opcoes[rotulos.index(escolha) - 1].nome
                if st.session_state.get("forcar_escolha_conta"):
                    if st.button("Voltar para identificação automática"):
                        st.session_state["forcar_escolha_conta"] = False
                        st.rerun()
else:
    st.info("Informe o balancete acima para o app identificar, a partir dele, qual conta banco será conciliada.")

st.divider()

# ---------------------------------------------------------------------------
# Execucao
# ---------------------------------------------------------------------------

pronto = all([extrato_path, balancete_path, empresa_codigo, conta_contabil])

if not pronto:
    st.info("Preencha o código da empresa, informe extrato + balancete (razão é opcional) e escolha a conta banco "
             "(a partir do balancete) — a conciliação roda automaticamente assim que tudo estiver pronto. "
             "Competência também é opcional: se deixada em branco, é detectada do balancete.")

if pronto:

    def _mtime(path: Optional[str]) -> float:
        return Path(path).stat().st_mtime if path else 0.0

    @st.cache_data(show_spinner=False)
    def _processar_cached(
        extrato_path, razao_path, balancete_path, conta_contabil, competencia,
        saldo_inicial_razao, saldo_inicial_extrato, dias_tolerancia,
        _extrato_mtime, _razao_mtime, _balancete_mtime,
    ):
        log = []
        _print_original = print
        import builtins
        def _capture_print(*a, **kw):
            log.append(" ".join(str(x) for x in a))
        builtins.print = _capture_print
        try:
            resultado = cb.processar(
                extrato_path=extrato_path,
                razao_path=razao_path,
                balancete_path=balancete_path,
                conta_contabil_nome=conta_contabil,
                competencia=competencia or None,
                saldo_inicial_razao=saldo_inicial_razao,
                saldo_inicial_extrato=saldo_inicial_extrato,
                dias_tolerancia=dias_tolerancia,
            )
        finally:
            builtins.print = _print_original
        return resultado, log

    out_dir = _tmpdir / "saida"
    out_dir.mkdir(exist_ok=True)

    with st.spinner("Conciliando automaticamente..."):
        try:
            # cache_data devolve uma COPIA a cada chamada - seguro mutar
            # `resultado` abaixo (aplicar_contrapartida_padrao) sem
            # corromper o que ficou em cache.
            resultado, log = _processar_cached(
                extrato_path, razao_path, balancete_path, conta_contabil, competencia,
                saldo_inicial_razao, saldo_inicial_extrato, int(dias_tolerancia),
                _mtime(extrato_path), _mtime(razao_path), _mtime(balancete_path),
            )
        except Exception as exc:
            st.error(f"Erro ao processar: {exc}")
            st.stop()

    if not razao_path:
        st.info("Razão não informada — todos os itens do extrato foram tratados como pendentes de lançamento "
                 "(sem uma razão prévia para comparar).")

    # -----------------------------------------------------------------
    # `resultado` vem de @st.cache_data (uma COPIA nova a cada rerun) -
    # entao as escolhas manuais (grupo filtrado / individual) de reruns
    # anteriores precisam ser reaplicadas aqui, ANTES de contar quantas
    # pendencias restam, senao o contador nunca desce e cada selecao
    # dispara um st.rerun() que reaplica de novo -> loop infinito.
    # -----------------------------------------------------------------
    def _chave_estavel_mov(idx: int, mov) -> str:
        return f"{idx}_{mov.data.isoformat()}_{round(mov.valor, 2)}_{hash(mov.descricao)}"

    _idx_por_mov = {id(m): i for i, m in enumerate(resultado.pendentes_banco)}
    _pend_manuais = st.session_state.setdefault("pend_contas_manuais", {})
    if _pend_manuais:
        for _idx_mov, _mov in enumerate(resultado.pendentes_banco):
            if _mov.status != "conta_nao_identificada":
                continue
            _codigo_salvo = _pend_manuais.get(_chave_estavel_mov(_idx_mov, _mov))
            if _codigo_salvo:
                cb.aplicar_contrapartida_padrao(
                    [_mov], resultado.conta_banco,
                    conta_saida_codigo=_codigo_salvo, conta_entrada_codigo=_codigo_salvo,
                )

    # -----------------------------------------------------------------
    # Correcao em massa: pendencias sem contrapartida especifica podem
    # receber uma conta padrao (ex.: Fornecedores / Adiantamento de
    # Clientes) aplicada de uma vez a TODAS as saidas/entradas nessa
    # situacao - assim elas entram no arquivo de importacao Dominio em
    # vez de ficarem de fora. O balancete continua sendo a fonte: so
    # aparecem contas que realmente existem nele.
    # -----------------------------------------------------------------
    pendentes_sem_conta = [m for m in resultado.pendentes_banco if m.status == "conta_nao_identificada"]
    qtd_saida = sum(1 for m in pendentes_sem_conta if m.valor < 0)
    qtd_entrada = sum(1 for m in pendentes_sem_conta if m.valor >= 0)
    conta_saida_codigo: Optional[str] = None
    conta_entrada_codigo: Optional[str] = None

    if pendentes_sem_conta:
        st.warning(
            f"{len(pendentes_sem_conta)} pendência(s) sem contrapartida específica identificada "
            f"({qtd_saida} saída(s), {qtd_entrada} entrada(s)) — sem uma conta escolhida abaixo, "
            "elas ficam de fora do arquivo de importação Domínio (nunca inventamos o código)."
        )
        rotulos_contas = ["(não atribuir)"] + [f"{c.codigo} — {c.nome}" for c in resultado.contas]
        csa, cse = st.columns(2)
        if qtd_saida:
            with csa:
                escolha = st.selectbox(
                    f"Conta padrão para as {qtd_saida} saída(s) sem fornecedor identificado",
                    rotulos_contas, key="sel_conta_saida_padrao",
                )
                if escolha != "(não atribuir)":
                    conta_saida_codigo = escolha.split(" — ")[0]
        if qtd_entrada:
            with cse:
                escolha = st.selectbox(
                    f"Conta padrão para as {qtd_entrada} entrada(s) sem identificação",
                    rotulos_contas, key="sel_conta_entrada_padrao",
                )
                if escolha != "(não atribuir)":
                    conta_entrada_codigo = escolha.split(" — ")[0]

        if conta_saida_codigo or conta_entrada_codigo:
            qtd_aplicados = cb.aplicar_contrapartida_padrao(
                resultado.pendentes_banco, resultado.conta_banco,
                conta_saida_codigo=conta_saida_codigo, conta_entrada_codigo=conta_entrada_codigo,
            )
            st.success(f"{qtd_aplicados} pendência(s) resolvida(s) com a conta padrão escolhida — "
                       "já entram no arquivo de importação Domínio abaixo.")

        # -------------------------------------------------------------
        # Filtro + tratativa agrupada ou individual: a conta padrao acima
        # so cobre "todas as saidas"/"todas as entradas" de uma vez - aqui
        # da pra filtrar por descricao (ex.: nome de um fornecedor) e
        # aplicar uma conta so aquele grupo filtrado, ou linha por linha.
        # -------------------------------------------------------------
        pendentes_restantes = [m for m in resultado.pendentes_banco if m.status == "conta_nao_identificada"]
        if pendentes_restantes:
            with st.expander(
                f"🔍 Filtrar e tratar pendências agrupadas ou individualmente "
                f"({len(pendentes_restantes)} restante(s))"
            ):
                filtro = st.text_input(
                    "Filtrar por descrição (ex.: nome do fornecedor/cliente)",
                    key="filtro_pendencias",
                )
                filtradas = (
                    [m for m in pendentes_restantes if filtro.strip().lower() in m.descricao.lower()]
                    if filtro.strip() else pendentes_restantes
                )
                st.caption(
                    f"{len(filtradas)} pendência(s) encontrada(s)"
                    + (" com esse filtro." if filtro.strip() else ".")
                )

                if filtradas:
                    st.markdown("**Aplicar uma conta a todas as pendências filtradas de uma vez (agrupado):**")
                    cg1, cg2 = st.columns([3, 1])
                    with cg1:
                        escolha_grupo = st.selectbox(
                            "Conta para o grupo filtrado", rotulos_contas, key="sel_conta_grupo_filtrado",
                        )
                    with cg2:
                        st.write("")
                        aplicar_grupo = st.button("Aplicar ao grupo")
                    if aplicar_grupo:
                        if escolha_grupo == "(não atribuir)":
                            st.error("Escolha uma conta antes de aplicar ao grupo.")
                        else:
                            codigo_grupo = escolha_grupo.split(" — ")[0]
                            for _mov_grupo in filtradas:
                                _chave = _chave_estavel_mov(_idx_por_mov[id(_mov_grupo)], _mov_grupo)
                                _pend_manuais[_chave] = codigo_grupo
                            qtd_grupo = cb.aplicar_contrapartida_padrao(
                                filtradas, resultado.conta_banco,
                                conta_saida_codigo=codigo_grupo, conta_entrada_codigo=codigo_grupo,
                            )
                            st.success(f"{qtd_grupo} pendência(s) do grupo filtrado resolvida(s).")
                            st.rerun()

                    st.markdown("**Ou tratar uma pendência por vez (individual):**")
                    # Renderizar um st.selectbox com TODO o plano de contas pra
                    # cada pendencia de uma vez trava a tela por minutos quando
                    # ha centenas delas - por isso so mostra ate 30 por pagina
                    # (nao um limite bloqueante: da pra folhear todas as
                    # pendencias, filtradas ou nao, 30 em 30).
                    _PAGINA_TAMANHO = 30
                    _total_paginas = max(1, -(-len(filtradas) // _PAGINA_TAMANHO))
                    _chave_pagina = "pagina_pendencias_individual"
                    if st.session_state.get(_chave_pagina, 1) > _total_paginas:
                        st.session_state[_chave_pagina] = _total_paginas
                    if _total_paginas > 1:
                        _pc1, _pc2 = st.columns([1, 3])
                        with _pc1:
                            _pagina_atual = st.number_input(
                                "Página", min_value=1, max_value=_total_paginas, step=1,
                                key=_chave_pagina,
                            )
                        with _pc2:
                            st.write("")
                            st.caption(f"{len(filtradas)} pendência(s) no total — {_PAGINA_TAMANHO} por página "
                                       f"({_total_paginas} página(s)).")
                    else:
                        _pagina_atual = 1
                    _inicio_pagina = (_pagina_atual - 1) * _PAGINA_TAMANHO
                    for _idx_ind, mov in enumerate(filtradas[_inicio_pagina:_inicio_pagina + _PAGINA_TAMANHO]):
                        chave_ind = (
                            f"sel_ci_{_inicio_pagina + _idx_ind}_"
                            f"{hash((mov.data, round(mov.valor, 2), mov.descricao))}"
                        )
                        ci1, ci2, ci3 = st.columns([2, 3, 3])
                        with ci1:
                            st.caption(f"{mov.data:%d/%m/%Y} · {cb.fmt_money(mov.valor)}")
                        with ci2:
                            st.caption(mov.descricao[:70])
                        with ci3:
                            escolha_ind = st.selectbox(
                                "Conta", rotulos_contas, key=chave_ind, label_visibility="collapsed",
                            )
                        if escolha_ind != "(não atribuir)":
                            codigo_ind = escolha_ind.split(" — ")[0]
                            _chave_ind_estavel = _chave_estavel_mov(_idx_por_mov[id(mov)], mov)
                            if _pend_manuais.get(_chave_ind_estavel) != codigo_ind:
                                _pend_manuais[_chave_ind_estavel] = codigo_ind
                                cb.aplicar_contrapartida_padrao(
                                    [mov], resultado.conta_banco,
                                    conta_saida_codigo=codigo_ind, conta_entrada_codigo=codigo_ind,
                                )
                                st.rerun()

    log_saidas = []
    _print_original = print
    import builtins
    def _capture_print(*a, **kw):
        log_saidas.append(" ".join(str(x) for x in a))
    builtins.print = _capture_print
    try:
        info = cb.gerar_saidas(str(out_dir), resultado, empresa_codigo=empresa_codigo, cod_historico=cod_historico)
    finally:
        builtins.print = _print_original

    # so grava no historico quando a combinacao de entradas mudar - sem
    # isso, qualquer clique na tela (Streamlit reexecuta o script inteiro
    # a cada interacao) duplicaria o lancamento no historico.
    _assinatura_execucao = (
        empresa_codigo, competencia, conta_contabil,
        _mtime(extrato_path), _mtime(razao_path), _mtime(balancete_path),
    )
    if st.session_state.get("historico_ultima_assinatura") != _assinatura_execucao:
        historico.registrar(
            escritorio_id=st.session_state.get("escritorio_id"),
            usuario=st.session_state.get("usuario_logado"),
            empresa_codigo=empresa_codigo,
            empresa_nome=empresa_nome_selecionado,
        )
        st.session_state["historico_ultima_assinatura"] = _assinatura_execucao

    if info.get("aviso_ocr"):
        qtd_alertas = info.get("qtd_alertas_valor", 0)
        detalhe_alertas = (
            f" **{qtd_alertas} lançamento(s)** ficaram marcados com ⚠️ no espelho abaixo — o valor "
            "lido não bateu com a variação do saldo impresso no extrato, então comece a conferência "
            "por eles." if qtd_alertas else
            " Nenhum lançamento ficou marcado como divergente do saldo impresso, mas ainda assim "
            "vale uma conferência rápida."
        )
        st.error(
            "⚠️ O extrato foi lido via **OCR** (reconhecimento de imagem), porque o PDF não tinha "
            "texto real embutido — comum em PDFs gerados por \"Imprimir em PDF\" do navegador. "
            "OCR pode errar dígitos em valores monetários." + detalhe_alertas
        )

    st.success(f"Conciliação concluída automaticamente — competência: {resultado.competencia}.")

    with st.expander("Log de execução"):
        st.code("\n".join(log + log_saidas))

    espelho_path = info["espelho_path"]
    memoria_path = info["memoria_path"]
    import_path = info["import_path"]

    espelho_texto = espelho_path.read_text(encoding="utf-8")
    diferenca_ok = "DIFERENCA: R$ 0,00" in espelho_texto or "DIFERENCA: -R$ 0,00" in espelho_texto
    if diferenca_ok:
        st.success("Saldo fechou com tolerância ZERO ✓")
    else:
        st.warning("Saldo NÃO fechou — revise as pendências abaixo antes de importar.")

    st.caption(f"Arquivo de importação Domínio: {info['qtd_prontos']} lançamento(s) pronto(s) "
               f"(de {len(resultado.pendentes_banco)} pendência(s) do extrato).")

    st.subheader("Espelho de conciliação")
    st.markdown(espelho_texto)

    st.subheader("Downloads")
    dl1, dl2, dl3, dl4 = st.columns(4)
    with dl1:
        st.download_button("⬇ Espelho (.md)", espelho_path.read_bytes(),
                            file_name=espelho_path.name, mime="text/markdown")
    with dl2:
        st.download_button("⬇ Memória (.csv)", memoria_path.read_bytes(),
                            file_name=memoria_path.name, mime="text/csv")
    with dl3:
        st.download_button("⬇ Importação Domínio (.txt)", import_path.read_bytes(),
                            file_name=import_path.name, mime="text/plain")
    with dl4:
        if st.button(
            "🆕 Novo lançamento", use_container_width=True, key="btn_novo_lancamento_downloads",
            help="Limpa extrato, balancete, competência e pendências deste lançamento, pra rodar o "
                 "próximo (mantém a empresa selecionada).",
        ):
            _resetar_lancamento()
            st.rerun()

    st.caption("Antes de importar no Domínio: confira o código de histórico contra a tabela do escritório. "
               "O .txt de importação segue sempre o leiaute de 10 colunas separadas por ';' — pendências "
               "marcadas 'A CONFIRMAR' no espelho acima nunca entram nesse arquivo com texto inválido; "
               "elas só entram quando resolvidas (individualmente ou via conta padrão em massa acima).")
