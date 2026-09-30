"""Leitura da planilha de empresas (.xlsx) do RPA NF GO.

Mesmo padrão de rpa/issweb/planilha.py: a planilha vira bytes numa linha de
rpa_execucoes e cada empresa vira linhas em rpa_empresas. Cada empresa gera
DUAS linhas de fila (obrigacao ENTRADA e SAIDA, nessa ordem): o portal da
SEFAZ-GO consulta um tipo de nota por vez (Entrada, depois "Nova consulta"
com a mesma Inscrição Estadual para Saída), e cada consulta tem seu próprio
ZIP, quantidade de notas e print de evidência.
"""

import re
from io import BytesIO

from openpyxl import load_workbook

COLUNA_CODIGO = "Código da Empresa"
COLUNA_IE = "Inscrição Estadual"
COLUNA_RAZAO_SOCIAL = "Razão Social"
# aceita os dois nomes: "CNPJ" (planilha nova) ou "CNPJ/CPF" (mesma planilha
# dos outros módulos do Hub, reaproveitada)
COLUNAS_CNPJ = ["CNPJ", "CNPJ/CPF"]
COLUNAS_OBRIGATORIAS = [COLUNA_CODIGO, COLUNA_IE]
TIPOS_NOTA = ["ENTRADA", "SAIDA"]


class PlanilhaInvalida(Exception):
    pass


def _mapear_cabecalho(ws) -> dict:
    return {str(c.value).strip(): i for i, c in enumerate(ws[1], start=1) if c.value}


def normalizar_ie(valor) -> str:
    """IE de Goiás tem 9 dígitos (ex.: 10.123.456-7). Excel costuma comer o
    zero à esquerda quando a célula é numérica, e o campo do portal aceita só
    dígitos - tira máscara e completa com zero à esquerda até 9."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    digitos = re.sub(r"\D", "", str(valor))
    if not digitos:
        return ""
    return digitos.zfill(9)


def ler_empresas(conteudo: bytes) -> list[dict]:
    """Retorna [{'codigo', 'cnpj_cpf', 'razao_social', 'inscricao_estadual',
    'obrigacao'}] - duas entradas por linha da planilha (ENTRADA e SAIDA).
    Linha sem Inscrição Estadual é ignorada (não dá pra consultar sem ela)."""
    try:
        wb = load_workbook(BytesIO(conteudo), data_only=True)
    except Exception as exc:
        raise PlanilhaInvalida(f"Não foi possível abrir o arquivo como planilha Excel (.xlsx): {exc}") from exc

    ws = wb.active
    mapa = _mapear_cabecalho(ws)

    faltando = [c for c in COLUNAS_OBRIGATORIAS if c not in mapa]
    if faltando:
        raise PlanilhaInvalida(
            f"Faltam colunas obrigatórias na planilha: {', '.join(faltando)}. "
            f"Colunas encontradas: {', '.join(mapa.keys()) or '(nenhuma)'}."
        )
    coluna_cnpj = next((c for c in COLUNAS_CNPJ if c in mapa), None)

    empresas = []
    for linha in range(2, ws.max_row + 1):
        codigo = ws.cell(row=linha, column=mapa[COLUNA_CODIGO]).value
        if codigo is None or not str(codigo).strip():
            continue
        if isinstance(codigo, float) and codigo.is_integer():
            codigo = int(codigo)

        ie = normalizar_ie(ws.cell(row=linha, column=mapa[COLUNA_IE]).value)
        if not ie:
            continue

        cnpj = ""
        if coluna_cnpj:
            cnpj = str(ws.cell(row=linha, column=mapa[coluna_cnpj]).value or "").strip()
        razao_social = ""
        if COLUNA_RAZAO_SOCIAL in mapa:
            razao_social = str(ws.cell(row=linha, column=mapa[COLUNA_RAZAO_SOCIAL]).value or "").strip()

        for tipo in TIPOS_NOTA:
            empresas.append({
                "codigo": str(codigo).strip(),
                "cnpj_cpf": cnpj,
                "razao_social": razao_social,
                "inscricao_estadual": ie,
                "obrigacao": tipo,
            })

    if not empresas:
        raise PlanilhaInvalida(
            "Nenhuma empresa com Código e Inscrição Estadual preenchidos encontrada na planilha."
        )
    return empresas
