"""Leitura da planilha de empresas (.xlsx) enviada pela tela do hub.

Diferente do script local (rpa-issweb/issweb/planilha.py), aqui não há
arquivo/estado de controle em disco — a planilha enviada vira bytes numa
linha de rpa_execucoes (auditoria/reprocesso) e cada empresa vira uma linha
em rpa_empresas (o Postgres é quem faz o papel de "planilha de controle"
daqui pra frente). Esta função só faz o parse inicial do upload.
"""

from io import BytesIO

from openpyxl import load_workbook

COLUNAS_OBRIGATORIAS = ["Código da Empresa", "CNPJ/CPF"]
COLUNA_MUNICIPIO = "Município"
MUNICIPIO_ALVO = "SENADOR CANEDO"
OBRIGACOES_PADRAO = ["DMS", "REST"]


class PlanilhaInvalida(Exception):
    pass


def _mapear_cabecalho(ws) -> dict:
    return {str(c.value).strip(): i for i, c in enumerate(ws[1], start=1) if c.value}


def ler_empresas(conteudo: bytes) -> list[dict]:
    """Retorna [{'codigo', 'cnpj_cpf', 'obrigacao'}] só das linhas cujo
    Município seja Senador Canedo (ou sem coluna Município — planilha
    dedicada só a esse portal). Levanta PlanilhaInvalida com mensagem clara
    se faltar alguma coluna obrigatória.

    A planilha lista cada empresa uma vez só (código + CNPJ/CPF) — o sistema
    processa DMS e REST automaticamente para cada uma (duas linhas de fila,
    uma por obrigação, geradas aqui a partir de uma linha só da planilha)."""
    try:
        wb = load_workbook(BytesIO(conteudo))
    except Exception as exc:
        raise PlanilhaInvalida(f"Não foi possível abrir o arquivo como planilha Excel (.xlsx): {exc}") from exc

    ws = wb.active
    mapa = _mapear_cabecalho(ws)

    faltando = [c for c in COLUNAS_OBRIGATORIAS if c not in mapa]
    if faltando:
        raise PlanilhaInvalida(
            f"Faltam colunas obrigatórias na planilha: {', '.join(faltando)}. "
            f"Colunas encontradas: {', '.join(mapa.keys()) or '(nenhuma)'}."
        )

    empresas = []
    for linha in range(2, ws.max_row + 1):
        codigo = ws.cell(row=linha, column=mapa["Código da Empresa"]).value
        if not codigo:
            continue

        if COLUNA_MUNICIPIO in mapa:
            municipio = str(ws.cell(row=linha, column=mapa[COLUNA_MUNICIPIO]).value or "").strip().upper()
            if municipio and municipio != MUNICIPIO_ALVO:
                continue

        codigo_str = str(codigo).strip()
        cnpj_cpf = str(ws.cell(row=linha, column=mapa["CNPJ/CPF"]).value or "").strip()
        for obrigacao in OBRIGACOES_PADRAO:
            empresas.append({"codigo": codigo_str, "cnpj_cpf": cnpj_cpf, "obrigacao": obrigacao})

    if not empresas:
        raise PlanilhaInvalida(
            "Nenhuma empresa de Senador Canedo encontrada na planilha "
            "(confira a coluna Município e se há linhas preenchidas)."
        )
    return empresas
